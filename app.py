# -*- coding: utf-8 -*-
"""
Dashboard de fármacos e insumos: estado de stock, proyección de demanda y alertas.
Construye y cruza automáticamente la base de datos de fármacos desde los archivos crudos
de 'reportes.zip' (Stock centralizado, Licitaciones, ICP Cenabast, Consumos por bodega,
Programación y Vencimientos) sin necesidad de subir Reporte_Farmacos(NUEVA).xlsx.
"""
import base64
import calendar
import glob
import io
import lzma
import os
import re
import zipfile
from datetime import date, timedelta

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

st.set_page_config(page_title="Fármacos e insumos: estado y demanda", page_icon="💊", layout="wide")

# ----------------------------------------------------------------------------
# Catálogo maestro semilla (685 fármacos, ZGEN e historia ENE-MAY 2026)
# ----------------------------------------------------------------------------
ARSENAL_SEED_B64 = (
    "/Td6WFoAAATm1rRGAgAhARwAAAAQz1jM4SBuUQJdADIYSu+yI2OKDiII6suHZwzGZUBIcJVhsoOxrmLkOm4wTy97ToNm5nz6DOwqeAhw8qg0oL7weKmBee6t"
    "l6DQEjsJRmkOYGPlzlNnTAexClYgXR050w0UcQqbrCDljXLdG43JJPQ8sIzOtetd6AGPyYhwJTtVliV0AogozMYvSifBUvMqg/M+o5QOs0gBzU4teJxT+9xo"
    "OG3qKrHWquNg+uX2PlS0WYG5EOZAoGHCa3HmLEAijwoDvWxpooHE3XpBLSm4BZg1ObasIDXvOKYjzjCBmuK9xLnCtl6gC7kqoy6uIPQO25l4Q4vC6kvN7I+1"
    "5skdq4gOWPqlndh2uPphYhJSDnDd9drZQg1i71GPluFrYCFeWyCMzVXnP/Z0chcZOOS2XK0Qkp9PJrO/E9J4014gmB5+bwVJuqENR0K0yyOSQdmEfjOZgFh2"
    "KVr2TwVxiq7jky5cMvi1XAy542WZ3HPDnTnCpdHn6wFOYZjNPH9y4LRuK8EvnouNJ6nYnm3ll74oiDGnjfXWkMOLMnBWDgED5VrOOE8OUNEzkQzMbmPpMEUO"
    "CeXrFXecHCB3yu96FjYs9cLfDC5QwtdQMCMoUn2YzrYwNhEA6h9+BCmCaKVNwiKVUFg7gOzJE/vrWBwVqwYlC/WLCaCtwoTgi4RrZJ0zv7KtxOfK8lF0KFl7"
    "yeg1XPh6BGyPyUMBhOjElHYmYeEnBxHYqh+tvij2OscPRoug2GyiISq8EZFTCrtIAn2siz/XxiBAb2SvS2vaMftoPccRIIXBN6Zu0GIfDI/o5a8zlW/+LPdj"
    "gRfk89PEi9Awl9+I3xzyJzhlF0j35uzvz2bXa+KafYs17lyuUeqg+mmBK6gziZuCItI/74KA1sWrkw8NQMb9qv9QiPyjuLuo/0dwluk3es4sEqqkCUguZYtG"
    "thNFhRayeOypyoSumbhKhO8XfOCCSO25LZMPNg4NBMjXcjk34r5CTuBRlp3CWwbWS47vK/atFOE8z6K/W2h7DlVkebNBCfnahIwL57mcTSS6IdR/O52K5wjR"
    "KgB1opN1LGjniwSmJge3FwGIZGvKfTRv8CUs/NYFqaCF0UHRsgCLgV4AXus3xxYPsth7T9uxrJmfZqpzrMioO9LxWYO4lHljPeiV/JmCalxHrWBaPunQlfXr"
    "VVLuLb57mMbwF1Es1YhHuW4VDLOYVm27N34pPCBdmdHidb9QWY05d68DpI09ulGTkWR4AHRhL3RNtf6WIDnL28hO6Y++nQkjjTUKGF+pORy+m485NZlpckhn"
    "ThhPWyL5JfzU/67RK/V+wv2bHcG67axCxiWP+DsBXZRpGC3W6vt5n2naYfZMR0PmBJgJVS6+gc0EbYF8wJPT5qbIlvkdK6OhNCfzGJX3DrwXWUxtADOtkn7+"
    "P1OjNNt01KMb+JceT7b/M3R1nhmnZwCWyfHOWmMFyLXwxOgL2M7HMte9irqjGZYTBoZ4tH6/ctZteO4pkTXrE0wP2oWrJNfiO1FxXWWz5vYv2nE2AAoYEBQJ"
    "f4o9Ktht6OV4WL7oD40RoZESr9HG+Fes9nWb6Ap8kOa6Mz/C8EQ28rRtTNYvxBRaSkZluFpAChQNSIRgYCsJgtrzWq9LlC154ux6dvNCiI5CRE/5UfvtNRCN"
    "3L/UrIx0N+CwODuh8uOt2nEkFqsM6Q1MZA1droet7FFdPkKDir8y/Voj2g+KumLfLCaZ7Yql1ONZdD7yWaw7qmhMDSZKBPdPXewE5kANBDr6onUfOFvkvdbT"
    "VlLMYXQGrkjgzx9F+CYlBYnMhKC5HPZ59yIgXHmO4en0Uzn2hjY+CvMQggUDc+0e67P+mCyXh/53MRx2a11Tq6XDjR8n2oGM/9o9Y3cfBP9e9ktmzz173Ros"
    "FBCJmdDkFSrnKxBexj1e7ptZLscBZ2t7qqETyJewT5uiM+mb53xdCkGhjwLQl7X1wIcLG2oLN/6qbXpIYji6U5ddRnlPMGuDC/zVUmkQeGJeKBLp6Dydpv9D"
    "Rw5Sl+DmyGJo/Hb+5SeyMi9VU30USQStSBZup+O4jvCzDvzYnuvBvpQbWyopes0KAcydcsdw/sNWRA8Kop/+VrQ4iE8Ci1nM3UyiyX5HCu2iR8/F11ee+hUw"
    "W64eFWTpOeGw0FiW1R3/Tyi+GwDRqwzJocw/g7Jt2/wpEbahCZ/KA0H+XOPFrFP966IYzkqIVSIKZDULq92ZJL0QETJIfVGjGVoSWQ+OGzpTpBw18LaTEeNb"
    "Vw20JDKJIczeSks4MoMUDiuz2YIXrAEIKT8+qDM9wNtL62Qluae2NnVB6rDHQjiQlyKon/9jOA7T4oHr9Yt283dG2JoBmck/kb1Dt0BmTq9xZq36GOdSXbv3"
    "PEYtwM5zkmhVhurhPTNXHO0+y8pWeLYj7/Jl2EQ2nR7LKH4O1QhJ10tz/0vDleCIC3c6spLeOc2vJNbDQM5CBfl40KJwBijL/u+78deEgHZBgMPzsQ1cEPJD"
    "dNt9rQ3QDUIFUJU90dLrSCDzyqstN24Hn5pq8ICDvTDxcfcnzT22gk4vmtyaEAW2lyeshjovdIgOQDh0U2VkC4TQ72arvYHcwcz3MhRZwEVeQ6mumONl+kYt"
    "8ykXrqHeKC5oLNPR2u/vubpDw/3IjYQ+PTUxh8dycgzSVB6lhZWpFHwznBV+RjKw320RdUvPrwNJUf3uX91zu5LEDzSZ8zEtZxpz3AiwQnB/zWw3CeB8ZFq6"
    "aaDyH/YpmateoCtcroGT6sgnpbb6aHLzqPJlK1pUMLcazA5U5AVnMgUYqxpnJshbiYB7VhdFjak/yAVbLwR9FA3v4r9ii4bZs86jYF5b9w+GQW83RWodWlDr"
    "ZtX0Pzmw16KacYJWy5Zwv/iVH3ycpmnLggAZedirQHVz5n8f8iKZbqlNj5Eb0v2HESvHrfnZrkftnnxiqulQmSKzdk3UUunv6xHrtCheGIrf4R4jZs/o4pM0"
    "Bp4PbyqNhhqmnEFALyBks6ltePeMMzfHJizwazhdaXAOFz/pXN674CFLa7qENLq3OV/uWtWko0h1TTiQYcW+djAuMyAxBDNsq5jmxXP+MFsoXeij9Ck5bH9G"
    "0U6fDASL+zHyYmAf4aj22FBsygT3JUaO3FdmdUthtSjxkHJK0gQKMCgZo3pC6Cb9tLrQ5MEbHzmH7hSgVmI4dPuBs2V58UOX93RrEe+k8b8FE4dbmFDtCN/t"
    "QHk1eIRLFseCE8w49YApLcfSGTq4LK+GRt8zTIj1sz2D5NC/XRudwX0P+1wmmLTCCtsRuKvZWbIg7GH6S7ZY/AgEy322ce03YNsTn47qRkyWcLPwgoIuJL9+"
    "uTP8qEQaqwXcXWEokL3CJKMOr4R9F7azScxymn4Gs/RTCXKNJR+YbsMP9hS35ZcXfA9/Jgxvm5pVuxaPT9P22lfYdrQqtTd/55sJVoKPOhalSeO0H3HKG1Ti"
    "ipTvwn7X9LT8mlE4Aj3GwaviWyFEROuTPVr0iiYQE+FJ8WTEjdKlhb+tuqsPVqancyB+Lh/M7p2c2GygtAb54ZDIjwL+2sajyEy8SfwhF2ubKAscyXazr+vE"
    "vdlvH320LkNCt6fku/i+lMS433fRgjv9M3L10TS6kngL0vP2pt6sAwUCUeMcpnmO2HwO37wG8PvEArXbgM4CzKUTrglON0gJ7CPsxa5Jy2CvLwVahflPkZ0W"
    "HILSobPMA/dJBnaz3sRPYYE3qayXt3YF6UtlufPY5ZMgfQLcvG6WgOsQi7DsBByNFPusLk+Zr7yBOnrUxfd6cNs2jEvr6y/Nficu73qCfaBy1g3nYhMBLIG8"
    "K6ZDjLUgmm+CKlrJBXM9L0bhOR8oxQX+s3n9KirNsiIgbcEHtdnK8ddTxhbUgbBpM/OhuLZYD8SZJ0nj8TjmtGpz+4M3nFH9tdjRGrJAKsa5KX5BAmhWOCVE"
    "ae+MvvJ9ISh0RiSzbVvn8s0V+H88BiPxJ4hkE0zx2rfsu4AUwNC6i66GplUr1np0ZHu7Y3Ivy1O34JfNmN3acGf2AwnOk6yaizou1ZZLRTBAGblawTXp19mI"
    "T89Teo0iLa/J86ZH231SQrcKv9PNhOQoHWAFJWISIin2LguqlPuPtcU6WPD5/ZURFZDuQVBr6qqSr6oa6o9MZD8UNHUtTZpOyKLBwHXYvqWQDHt7wARJJYe1"
    "eaQyYkT0B/cEm34rEiCAYeEp2ejQfB1yDb8QVA6/CkbYf1BftwTlVb42JbaR+DwL2zTlFFHLM1NbK8jUYZLGfcSCpyBBWU7v4ty/CmjNli6P45o+1JrT2DY2"
    "I06L5MwcXtkZ3eHKuNp56byG7YIB1fglx2lEZuaC1aD1hqrnwrxopWBiT2ThnvmPSC+1od8nsVH1FbgRL1ufuODgo7WFe7ew2Zi2elAPuGxu96nMb8ssu/hr"
    "ihdIaGD++IlrAjpKo3rEnXRizPOu5OhhQxnAvcbB8KxY82DBL6S5FfRzUmyOIglu1I2isW+zccXqGlVhHg6lZvfoHGye0oyzvzaEGFudrBwqUzyGnIpR+f2S"
    "Kt9vlwReLo0cerstDww+ZEV7HLXtWxE1+ry7yJfju0Td6hDwyeJYmsMo5knc4kntUjOERh+rwG442Dk+A6JAD5tPfOP/YWAWZ0IPAKFJeL14nXstZN2hEh0a"
    "v3iGGUDtP0yPeauzXmkZb72YldUQEbX88SgCtYhz9vC1rshiF21KI205IbY/+BG8yXPh9AbiOUNit/Bkcf+10HCC9r6+nGFzIrsnf25MUdydrZ8Up0JkkrG0"
    "2eV99WpHQe67oXa59VIpx1A/zFC3qnhyI9KOnkjK5Y4xoHbfMJQIEJmCU4h6x7jQe5h1Vfi6vdT8imTfAO+CgFcv30HqyQvcGfqYSZ6WZqb/FntbmMLxV2nG"
    "aVS6AfoZGgCG0XTArjWR/atnK7lVrT6stv91yUKYPQDczWQV0rstRsalEyU/GhXfz6H9Q3zfO9lq8++DO2YLzSRvA3jbrlQC+4Zs1suSH4+x2pXzx4kp61bE"
    "Vju0UCd2eq3sIa52r0SXse/EMUcX4Ss1oAbV4VFs7nw4t8VokK2brWSz4MAuo5LLblUpjmMlLmYHbZ0O+zcPWhCLcPakxR9ibBxW06Z1Ct2JCcFpYR4A+lEA"
    "SFPDGdSlPfCipJAxhqKwyY4JGhW9YVeoiX7zAsnIjlYh4PA71kkjvZaBZB6CtnHrelCb/O1kGVhY6bE1PnlqZO8JetV8bUDBYC81Kxvsp1TLZjW+5mOuK883"
    "W3Irs7r2FlaVEL+LH9yqPFP66W9JAU4P1YJOV8EIwn1m9lYQ7nWT4hz7uP8ZR+9sJrnuGzmyxEBlcMSPBihHUwWnFpLPea+CS3oDtg1K4J8/JWMM8GTFAy9O"
    "fyewry58GaCZq9Szzxkq2z0yrC81xAA+kbRzbfXXs0rn86n/00ZfJNtwTQnxu1iRvXvfrClKCT3gI1F+RdwEiF7eRdAtWwBReb4e22DAmR3hABnC2iKGdfW4"
    "sIKx/pSKkMJX3nagCdyYeXvxHpebE0Rt7NIM0a5X752Iot1j0w44lRBeqfYRyVwZic9t+XsNxxF1hrndPk9exV4HZKCKsSAW1s7bAnDJz2QSBlg8pPDeBLaj"
    "yTKMQ7zheDs8f+yJVsMxBJDavifG3qFNi+qrGhEvu7hRZkd/CDttuxqIU6ls7f0tFP+EnWH4ekXZOytrD0uIJ3SoIrSl/wBzNjU7quRwW9f03LjYb8ADcNZM"
    "Ci9/LsKXW/XBz1AEfVaE2PcSeJgIFYim5jVKq2CYk6bOY/GtTp92v3hQ3FnCNXFACQUuxunr2n+63MiH/9Zg4+smbgAtpsbI897Cjz9gWog3qrT7nSNeStyn"
    "i8cutzF+ESJaTFTl0x8OUteYRQHK8EE+TR85Gxz5irwlxYRVP8l2AlS0EkkKmDiZba5bredv4LgvbvyyH4EYuBhX4sqTK7DhFq8HE15F0o2GSjvOQ2IwGn8A"
    "ZpeKhqdYOtkSL/j7/uwRSD4YeyLMIOLECJh6XRbZJhg1yAIq4rCgOBLL3bHHcI0CpSkVYb5PHg1WDkE5ofLlB8RfWD70iORWZLvsK4pAenMOMXuU/IRZ4E0d"
    "wRUDYJqhE494D0k6ZHifXDU8lOKqq1lH2hx5oeg2qCS/9FECtnvsRZ0HV93ktaP8XIWi9SL0bCEGG9dypFRHujUfBn259QwQ3EsldAoO6n7p8/v5Mg0wfcgm"
    "uwQHN5EIVrWUBBa5gY5DXQvS1FdwmVS2vMk5H05NmA6wUUD+Ddpk9cN/QFM9nKDGncpzJ0+35tZkuw1/o85uMCkQ6UP0sg83XN+T7h/fAqIcVoLZGXtb4Vqp"
    "405cooY3Gu8FtyMn/qJj4NjjYyO2lOtXWvG+5cmNb5CKHc1A1Sg7GEoGZzqrJ6pP/bE7AbotUDeGatKFjHSZ4B2I2F6+V1CWlOUpfmBYau1sH3oWqDJRYmP5"
    "MBeF0YPP34ZEyXelrengW/gYtRWCGXesyr66Rp+xVPLXJs8Q6IGzyquy6lnde4HylmyBvBoYpd9QeSqJz/eNTcd257drST0Vide8518fwx0qXyXBvbsKIyse"
    "6KVSCVCGf050AUMZOWk9D6tBFPguKsvMmw1GSat/ge3AoWpVP1NJAzlKae4chror9k8P0k6XM6HAZlpHsGTJz0OOBc4cZ8GKNMkHRWEOadDLLl4lBY2ziUBq"
    "3p+/KDzBuWvBb5dWh5kGlXcGwrnNLB4P9Emv2Cj8Emcd1KBIriw5x7zPo1MY0ukSX+Pqua/ygg7Hg6kVqhZZk+3wVgv2SX3g3zTDghlOJIPrODHh/wXe/qsT"
    "R9jJNasz8/Ny5HBxOhJ1jMxoWDWMoA8F/ObsQDKosMcTxYDqPtnr+wq4DPIdaNtbVmS7T90DBbpvo4GFgjZ3jWWUuDSttTzXsbWwHzWQSwdLIMmFjqmBHdno"
    "mBV5o/DBVdHbuMnKWRzbZ53TxJzs0WtODGScpmLzKINygh+44oNXzgrz4kVS3IvKVwUGzM2tvUMWhObUUnQm9DUGa8B2ftbC26lBfR+bAKPN/V+Z12KOJrwr"
    "bnczemtFmCb+f71p67cVb3wDnOucilLbsg7WwaDIdkKjsVd3QJ3Y7zw8iIQpVhFk/JISzVKlyO7QlMhgzvnjugJlW7fbpQfhdPehxLEW+20H1A0bwIVwMMJ6"
    "SrrZ8yP3yRQwaFm+B8VRdV/0k9LO6HLIJyvrTZouVrzrp5flf6DCPeOpNG8VDmERry+QPwb0L0fze4RWNQwxld8a8fM/TEyBjT8l2OyOJ5Q8fdAC3VXIPm3Z"
    "kMKI8GrjntKwYj8Y1mlUzU+RSsRIRTYdcJ1/0EflsxfZzhWjwTMLvRcbEbrWWckqjuhHFuLVgn5pb6IMO+8T1OiMTJLJ0if4HzlGwAbYpT7OMP1SJXs0X2Mk"
    "DaXOMpf9SDhaDupGP0OruYEHcHqaA2ynUYje8HfwQYvd7FzLQ76ZIKuGkEFzK36wggSPKveTjhn3Byo3TVlt5L+pAen7N6R9U12hNrTfYIPsod17scHaGHG0"
    "yU6COMHaVu/E2gTR/mNAIR3uihLJT0E5poEgNwAyYJYpystnZ7kkscDfd1blWKfodFUtZZaYXj/qrtXeGLuGUPTulbcZdKo2NNPjyQ9f7h3/g6r5phN98MBq"
    "P2380JZAHQpGPGrhPXqOEWRxuge4XOJj1rrSDaI6KY4mK0v0/dY/X2eaDCK4o+hJnkuBd/U7kppb1hvE30W6kH+I4qIUm6htdLUIoqB9wnfIQnlGUby3sTWE"
    "ELR21vz2RME+6rPkAV44dxy0+VKJPao2GuI8RmuJCFuJK0BEk1H4BRIjLiHPvt+Ekafqw2etaQtJcElsaIpHa3Qa2VVv34m4+LpSv51RDm2LPSgBqfGleCIH"
    "ZP5AVqOifDPqL3m7hs8YvwoO8KuhOUksRGhbDcnLALr36+NsvH1m4k17zNrmxmoM56JNL3s9HSkZFGH9L1ckwjbBmiVzPIMW1aXuV38nHoaA6PkaIkpg+rNn"
    "iFvRUUHEf/6TQOyUNwIJ90nKAC9hnjtirbBJOgyal4i7f4LKjHSn9J7E+C3tAHtdK69UN4ELk3PXMzthMPTUGfQCcGnkOn93N1/yD083YG33cqsjPAGUFfFe"
    "OHqIdhU6JZLVHiyynkGZQBiRv5MhNhZbzCpnjMUBc83GcRoehHgtvYHGifIST5dTQFiTN6Tijb5KQl1trZDrmdQo1o7o0RjgC9Wz6al1N4/ApxquwS7EQwsX"
    "kno4e+CMHCtXTfvfffqK6yqJCA0+JGcaa6RlMmYPDo+00cf0gpoLBexkR6NHf+Ea/AgjnpSQ8WKyprAsTV1RCdMFQgH9CXaz/kb4AwxaSCWcAVCooY6EXxHy"
    "3bXzcXHaqyIBkjmteBc5e/Q3JEr2buSYOGTOfb/w1YUraHCNs8BZmuSjKlQ3fv0EpaCiRODXMSH12uN0w+49+dB/ft7jNXsIV6Nx4+AMZdhUU5wVWg3u49pn"
    "kmZRbgH5Uk0A0GosufNElwfaT+BkN4QAZ5IMMQBDYWTGDUMyVjPRqWo96OfXfDGVwviM4SagIlaqNWEEyLV1CZekMoR1knwXgANtm4j4Lbka3Ye3nuN7Hb2a"
    "c0+xxiWwINIQdUiWNWbc4QfReQSXUtIEgMoeHgbwPknrkGETQSJWJkbVAMnMLPNu0BFEsIyFSsCiyT0e418oirh1NKeJJ0Dot1MWk5xum21DNenpaNHDJiA7"
    "KdErrRAk2FNjOk2z25lCV3LFEYqc5fcyxYD/EUA3UOzkh23xd042mkBXcSN8ZtwgCp1DDCbU5aIGm6b/dHdb2/zYqLckjTryRUp9fWy1e0fFmov7FVQu5qw7"
    "xyRxaEEOyy0J0MYbKDvpYKYsYxcGtcEdSJS3U2XE4fzDwCdOUb9mDDnFytEHR97ThpJJaFfqRJtKosK+MP3hd1JfNdAmT0yQk1HIWz4xFU8Hn2bsDFxw4Yxw"
    "VO1dSCDl4OjO4yWZO7L52GomFAUw389SAkl5Af2timXbzGEUIpbaVNNjoteKrWQjxxm84362hXd+zhcgu626uiO81JGa1YfztmFjtUTWQCm359ThNdEoWM3/"
    "GhvLYv2MPd2bIvEQXhs+bmDfyWkYTYKGWzRykzd3WdvP0CrBSB5xgmN5s1WoyHdfGeulpx96urdzhV6NR/6tcogMP4kDG8dWE1jysM72sRyZpWRsVC5ANMzk"
    "Ihj6JGgRlj3EWtp84dfGh/38DI5g31VWe99CnLk/rC901MOv9xKiR7KHgxqBMBUi2+/hfmMJdEP6v+A3MtHMDVUCXq1uYThrHfz0cv7V6qM/2/5Gd41mnTlT"
    "vxzcB8Y54xIbAUApEBANyIYOZqLVXTPsc5Rg85MxFPv1Ufj372T1yuQpInWm4AddU8rgSkpvn6wUqKc8snpXYgM1Rx3654qvnRE8Z9pdxBnIDs2nbHsbxQu7"
    "cV9mYBXoTXxCr50mGUuo74/840ywhWkKDyrcqySKdiDiIA+A0GTTwb/1Kofl56nTMg7PNY4zEDoJggjE8FkgbDGQb877vchcbiBZI6ElevhgGYWGK6nHeiju"
    "nOLpBGCko5oIvuYWsFXWSudQGLAtmd4s2sd2l+Eni9k+ctGDUlWya/9Zfd4jb4Als7Pol9GL9Ntapu8zRjF9bCv/UNgpQ5YuwtMIkk7srYrm7QGiUBm39NLK"
    "0JLjQWzuq2tXZMb+js8B8YuQhvWZ2Hrbr7rNITPXLf24Mfwj9sImHJhAmkexzP3O4N86v5hCdcJUyblTPldNHZVxUMcnCYMJEDth4C6/CXgPWm1Hm56MPGje"
    "IOzaUc9GeU68KG2BSyWy58dBvlfEBY+K9rqDezxjd0cSWxxuUlycHnox0uOqiDykVdG7InJ6Nb8U79Dnx7ACNByXPGE40zSf34E3Z5rORbkP3qyxCmrecDQf"
    "ADeIlKPMYuRuvV5PvQoQjAGcV5Jo6WpIiIzJBZYWqWEId8xFyEkEYe/lDg45ZzMwp3j1ymUXeCNGBx46bjQcZd8vQKvTWLz6QD+NQGk2SdslSWKHBqYwlSU/"
    "XzcAjN5vIXT+nuU/kuOESw2V+TA30qn1P1ISkV+Gq8GGxZaxWQqKKU7zMXo92C15i+K9t1n0vX3pw9s2SH65g9UL9Qmh7NRuUvdFpVNUEffG0fyQmmuMje+E"
    "lEb/+2lk49dcq7i6J5BjpxGhaEa/Vmv5AvR9tWF/1pl2HWv2Cv9g7dF0SM2cBmS5O1VJlLPe7E2YcL0hKxmehySROecfawyoKlTF5RNOVZ6mZs2tKNhTGENO"
    "wx7kGnV0UYt4ZtiE/KQaQtp0bnfOcHPNj/6jUKUZ7npmVk+fQvlNJVoJfaERSd0GX6BiV4OhjpX4MdkSdoH1ENo9mxKv14xQ12RPFlc6jtx+ISSCK1VNE5Ht"
    "qE7rwNbfmAIp/0K/k1UUHTy222FO9e4mYB2dHqgEhB+ByrB9AZgzY5HNUCOwGefch0Q+n1R6eBn3VZ3GiocSk8U/2ck5CQQOXEBNsesJ9FtkHL3KhLIWXuz7"
    "2cyvBae+m4dKak5b40GMILLsQO5jecory5DncK+5kiZOjnEaEYBsuP1Nf/VVXG1NBKheSfhQXbVoC70TxTlU1lzO5Gn5QLVTCiVlHQjRSBzN92JVhQfq+oNv"
    "sz1inXIhikBFkDHhgMZ34nll5jsd7lgOLAUhd2gaMwTpsnbqfmFJEf8aVNVZ6oR6pzrH0Kp5dhk4eXvaw0xL+Rw7zTJO1NBWiHQxsfZAHbdNX+pfYCim9LAS"
    "6TUnmpcYuq1TOnZY+TDT49/zLk1XYyGMy1gPeok/IuIh1DRKRUyjp8ADaW5ZEAvUJyVUtWKbUN0q34f//8BtyqTVMoMpkKXqIysG7bvq2zqsPgTpTnmA7qgd"
    "uevzw+7hpiY5URAYCwMrO/Mq978PR+E61xvGe8vNJP2C0Lj7OI41eXvQDUSFKvs+MKXaNM2/9fywAvd8K0S45fAlz56SIcTPLh1OQfM0G6hRRYUfipvg3Ia4"
    "E7ovI4RTtIdKYWVM2eCB8UNYQaoxsHY3FIuF3gZVMUx0LKuiH4Ya/vw/kMTbNo/eNgKbVn7VWkrX+4L9yDGd212nQmtXPOL4Aqrw94+WjM3hydk3+kZvXdes"
    "HkJ5JzeNKJiaOhO7LH23whEqJqv+SMTgI0MOHM7E4xX+CaJou7y+/fqAqQU/jiXeCifc527TRIVqPDY0EIgXEAhpmx2vNiVk4/ixL9VPmSPYNBQ7/6LOdQ8B"
    "qjbUUw/2/3VZHuVUwq+wQCeY9NpBq49FAJeO3KjJmukzbbvkh24c0Z6UkY5IWWZy88qAQ1BRlMOR4Q6tCm9ZENCy8w5Eb7o3HXabZvJDzbXIE0/1oIkvNlaK"
    "UV3JfrHdi2k/hr25dzwnxbrBZRGZ4bkTGv1OYj9LnIDjTrVfliDAVP9zEGwM/uWSqa2BnHJ1X7bAFt6H1LPoJGTPsfwoCTzzFp1UrxgpNWTObY9kF5YKB33I"
    "1C7MFDHv7qirqnfk1AAz25nBi2TEVkJlCdmZRBcKe50szRkR/tin++R3s2lAEaFWR5aLHKrPK/vc9M7pFvcY/WN6xPyFqUlp1KazlSBCV3RsyW6CLMz6i66N"
    "FVCbwwqZRDa7A21J+++o34vrwKa1/1KPeJt7O5l9jIGtSXbY3s+sEkCr7fLmeaYVQcM7DKh0Y0oHgkIWGfgXKmeWImyrXVkS8E6QTGc6aqClunnGN3GlDYMa"
    "rjw+6z48259tmMCCfgyEgtk99KTgLw6NvuLewwW7Gks6E51Y9TCfOc5CJ8mWUZ7p9+MM3GhfRh3REyHGG3EFqtk2zOK9QV/VN7GWBdo+6XEDSmofXvUE5C+r"
    "kjUiUcWzz2yEmm5QVywa3IG1xhLCyzU75NQomNuPNZW6x2FHJH5TRpnVyyvb0lcQ8AMwgPBMHv0qFkFMGCvt1HDrvxl5G8HRhbCAId444Y2d8r9PEgKoBMGe"
    "n2wYpPPLt0XA//6jZp/CQVSVi3I9VvgPQ9xvZkkUay0DLvoY7FVTLEpc9ydvdfsK+VTm4dYkxQwvH40Eu4Wn9nOVrSlY2cGF8b6Gmpe31hp9hd5cqMFb6N5Q"
    "nuxBJ5OwWgrQt8j2VQ37OTdGjBUTRYG7sHj71J64xBaRB+m2TRRmLUfUMpl7SP9EzjflcC8WCfi9/UTMLZvnTNNdlkf2X0WEM9wFos81tWO/7soI22u8oe5E"
    "T1l+Kysi4MdCMg86rMSunKF6+8IFD8CGmcf0XYhC1baNpR85aU9K69pQyPhtfH+t1BAUo0o8MffcVkbHFnZI5fuea/UaFIFGfrsYUDNNgtHJ1otNghUwuXSr"
    "k4DmPxyj9PsfBiBBsPDht0wvLIaR3mEphL5qaDYAClny2z4Id8SDA88V8MT6be19sWlwu+ns+NqdF/5ZUiDPvIPpfL1g7ZJ2iE2MxbqvY+Xs7KF0sucb+/Fj"
    "+cpEsows8HGnXoBf0eFOeG4ja0vRlAKQV5J0Ub4c21/Jtk2kZjE0PcJUpxlfSHF1FIl8fwKHSfTu1NH/XAbvuU3BRFAlqcS4mcnpXntP/7W9y/CiPaRMOBPc"
    "Bn/c6QCr9u4amX4Lof+SMidFtjhQ0ORm+XQ/Ls/ZAuEkX4IIEP9HpDvAZ4VEjB/YUoy7Hy2/byeTIoLBpUgvzrExB5wI/7VfGCemVcpMyp4gy1JTCSojIM9y"
    "uUC/mWgJMbfrdbBm4Xhsx2fqaIQvb6Ot0HTVbKLfx1OWjptpJfeqe2Hz9VWzoygt1VgArfzKbaXycnEp2/3UcthecOPuhA7eLPKszf5/8RaOq5CMfqSV3ZMa"
    "z7OGIUI/7K2cFr2KX0qVSbxZjDkZ/VM+2m88N9NhDmql+JDhWhj/GkMLNs5DPibjZLtLkVyAJ8oxsJJ28l89E7GYMaSHKQCxTksKbcTn44OnE3BbXjTxJcBI"
    "5WisZa0p+/qx1lIaqBGXrfMT0vfjQa+FbUqDO83NaytCIEoATrl7dEWUZeS+IV6W4VAwxuKQ2DkLuRBC4qhGmViu8dANciKqqljJ5VDTaHYv/qIQs9vcQGz6"
    "XsESgjxGWzO/XatKRMIQogq5ypKx6tEYuA66YSlJatcysi2udU9x0bLHG80rDuasLPt+UZxRZAHO4QjOOJiYgHlK/1ECTzEIkrI6lTMHxUQw2VhKlUEbrLFU"
    "AURPwQAjdyXhyMwCzjIZECUfpMq3NCxLzpA2dnE3hJqFCgWb07uubpK1k8uPC6R14hxj6dc+Ly6xjGDQiwndiPPxXjC0CPqt8547t6Fk87lkZY5gERt/vDE1"
    "z17Fwl2kg5vdSxwo9rIozG3oBy0+CQ+agJSszN+HFO80SZKCy5rDgmLq9Bd+5vt1X+62lImLHGr4oZcFvcIm3TWGr++Sp3sMTlfxDHt3d9+Rjb7m5zI6qrXL"
    "s6zSA7x4ooQTbxwnyUrQlmPYo0QNob/7J/fnLU2S3Kf48UC7vfFYloTRePtIC3ouQ39eeFXAVNwDWBC6LSUZV2Tqjn+gCuhdCrxDO93krY3O2zLJHo6iKIyL"
    "OsL4igqP4/jdAGkjfX0smSPc5bVZJfR+Pa4ejBIzE/9RGv4lZk2lkNK2F8LM5bCcZWc3XoxUVQeeOHcnhejqYE9HjOFiCro31ivgMZs4PeVTVR+Hg6cluGt1"
    "VDclkcusfMnrXNeQT0drQXuiAgKkWf4cLQ5sUhIYYyTVpl0qOyTJlOVESG5xhtmoPefijSY8y/0NZTk0xh4bQDCQcdeFML+tgAr/7FVu1Dvkxzuexeu2KslH"
    "VNH1AET6hLnj14bz/lzid9CF9BztbOsTsYD5ZfNWusFUjqgv5CNhe3i7Th8N8wffcH3nnILFvxV4LrE9PCOl3DMnj7/n/Nh9uEpGa+cVL810o06I1gIs4W00"
    "cPVTprhPLNuAFWu26tys69HPqaPdtsxL/C9z5nl4t6OshXSNPyldRdI8olq9a/x2haIq5lLBhpGl1/n50jNoQcuNexJVwBGCeFMkDZwtk/BkX5mjhaS2uGpP"
    "5d2ZfFRC4z4uq0JBEYupbadU3WQ3WBiWQrv3TDQddkNenROZ0W34luTUmetFS/QblFN6/0agZc/KGPku3uVgAxBOsd2K95gTFk2+pVhWVorTY25SdUZPKLTz"
    "owWDHQ1RaWJiJXcMGAZoDIjxwUqPLlHXZWxrFiUhlNOr9pQzbSuXXs3+1+Y5LDE6D/lvAVl5B4JmpB3ruXLJIQwyuaVOomXXZOPWhg8aLtITPy9K5LG1/bL4"
    "RlIpF/rMyKTBVayFnsdIM+VQt3FTGYjthbA7W6VguQGyJTsC15l4DxcWyENcKedEGGAnB+dgfOIUVcY4Ib0o43mLtDF+1V5t+v8Wp56S+0Cqkh1bII4ASY09"
    "cMLuhjDuJyCnaKK+rj0bNW3uSwEY9s8kLGZm/8MCbnHyjxfkEfFFCYJ2ILTTsHexDiJMNGJDpJLLOmrd6hFYTBRVyUMxCCxiMiHPDpecRF6AK80BNhSfEOXL"
    "HwzkcwrBeA8H/zz+LHMkHVv0VsASMATuHKyeS1d7thcK/UGYoII9LeBm2QV9p6ZbF4NgVSjLbf1n3woBvTiJG2cKAR+v6dG/L6iYmVyg+AhVTCaNhq8qDkc6"
    "FBbxnxSJy6j7bTDr6ZkRxYHqMRBHnwowadMVITRMExdJa9mBbSGeG7Y4RtEv448a1vZcAVe8nYJifYFYzyK0oGykNwrQihDAPQYNVnb3jp2Lr8OnF1wqZanm"
    "ZOmSOUsqEA5n6UuUsj3lmKK5S/sO6OB5tLsFPMDnDYzbiChFDZvwuD8pJm8Tcko89UG6FV3N/vEkQPcouGLGSmE2j0jZD3YKj9GBHmi/sV/Z+BG9xNW3HfQX"
    "IyrnXKuv9qtUePV5n6qhjDaZfaDSJoKFw1ZCqTeRB5JkG7iOrD9hF6oJxypD42mu4seKmbK1t0r8bfbB6Ghj51c54LX8JMhxzTp4+YRlQpAw1UlyCAknN9RQ"
    "Q+L1IQJvKPKHLvk2XjGYnL5D96J0YBJZH9un31zYYk2wM52eip73kqZhwHDAX5NAPJvSfeE7TwNsrhp7r3yAlasPqeM8MypaEzXErI9FxyrKaDGDxb6aFOCa"
    "rRJpmX4TyFoxmfzCdcL6pvVlCbpMpEPY5AEkjFlmCcqS5vbCp5qYBty1jz+2HQYtzCInk5wtEjCYWyZHfOTj+KF7lG8UW+qbzPiXkJJYhWX6+UzDQAHhczrd"
    "lu9Yf2BiFFiT4mLr0vwpJoyhOIFb75jVX2FWEWAKtqRDjVIYvobasYmDLPndhepNxw5I+Pf0OuKxnFkfz9af+nZU0hXwIlga+CFcQaJt1J3no3EHfchlg3pQ"
    "z8BxhOxjs9EpG2sL4lkYM+7x6aXcjDZvhH2yo8/CzCST7vqC/eEBvaYUC/9AXkRYUJKPNOXKhVxn3fYxJgzvLrbfvAePj1aQyi2uK5L2XOIKuHCj4i2qFsZg"
    "jpCXP5MY6lSY0RpkFl3/b2pyst9r3li0MV/7OwXV+7XeCOEoQ5AibUCERoaqkSUBcRHUKmj7MtPhxUlkB8ig/YhySr3tL0DQ8m1XCI/lQcU0Qz08tY7+2s08"
    "VHe6liNLIMPpF+OLp7hk6b8/RnbDalZ3KEuTW7CDDzrK6ldWbRq4ea3qAb20TAriVp+SHpsfOhblHlqIa4ABPS8BB32uzDtdBeRA5P8/+n+muyZhSAUavF0l"
    "HzpduiGvYqY3+R+RLkCEKrd1+5xYIt7Iby7K9AAxnod77EJr17hNhRzTuot1soRAqVe6+iVjF/WZ1jMu0BN2o91/wa1nrP9FaM2Zwvv9yubGVRBV/a3BvANx"
    "vO7jFwlqNMiiBEJyzOVTcNHT9SA/4TroTuKNCtfGDp91iId+FSLuuPGZ/Cla7ygq0vdF+tdCjsf/aX/WRq12nOgZD9DX8DP9NZhOOi5PceWhe65VyH2GHEPz"
    "DvfLUEO+pmrfRc/EfMmmb9Y0ThRNYaWtTEtUO3oYn9yFOgSbwSIRMxDEoECFmDtzZ6oOi1X10b8xLvQfDjYTpvHrC4DWcAfe9HaiYe6zr9Glivhj4U0AEeyu"
    "cGV3oeBo8uCuyr6DBiJYP1nq4TU0vgAhR2dODQUjAzppWUg2697XVKlmHGlRKuWP5vDu/LC+D/KmvmLORNPLNTpx/va75aF9Y5ORsEQOa/r28knSpQ2fg1B/"
    "/7nnt+r9VQya68jaOhOJQcWjXHc9aEPtHD2zwPwPlPmBk66EcL1ypuTKKzmyPc4YVgOWqJmipWPsy91NFMjAwforWEEDVLHNNp8l7y8A7kowhlF+ki+mi7uK"
    "sh8bafdZ8I5gBpOLc4h4lIO2/KDbfQJwuNGtrFC14d/y+gRawnVTh2KUv0CVhWq23fXGSQvnV4qf8FK7peMT3TC52nZNLKfGdF+nf2OinMtarZz0nwTwBTtj"
    "nEf/aPRwmchFWiVDAkWwkhKSdSzRYJqGepd94uIV5Qm0hh65u5JG+MDc14bTkB4HVSpYh//L7u6vYykZ5r9HihdXriMQvhwlhGXkPx2WudjWmoYmwIbal8S/"
    "XJiBjGn1HOg9mjqVstE8UNxH9khVp+MaEGTbWOZ+FzjDTWKzvlAOSFhRQS643M5QTxb8IxIydhA3IIGZM8KsCgvmjzK1TuPWygcUUzU/nxrv3pUyW+WmVFjV"
    "+fch3JgG+24P+C2bYP7I4MF+ejS79kqT5FeDBOwqaTxKrFuhwXLgnPHz8uT/S2+LrBrrpiV+/99Y8F6XFvqT6CV1zudvqEk29B5wrPrh72QGFIWTzyvXtaWi"
    "/k9TQ2ZJU217GwrC7yjILhrXnQLAwaa27WJ9VD+2EIG259epckn75gvJyLvr5qNjCNQERsdOeT5/7y5Y1LyKr5yqI/5n7/BucVV9t+hNqGpahPQ2+S+taHH3"
    "PeTQJV4MwuFW4JHP9ozWKO0P4WSNKcDnzRjQzNP0jchz8XXNzhRcVoSvf8O+a4x0o8pusmeE0sXpRcjflOME1oJEqDr2xMP7KVpEFyKMtlkswNvh4xAb/jXY"
    "pvnWwv60XQUKzLyXdp2QVrukok+lCMWmncI4AO0dvbvUYO3K3ZnE0AG/l6vyyc/WPwZXDoPiAilgbZ2c/vBmaS2K5HBFNZCG5YvToMT3iJwg75859nYBpSK4"
    "PepoDMWTlStXZ63JEK4iXCKsV9cwdy7dyyCzi5IjrrBAKqFCXNH3uEJB5duz81LaSApFNjTd3FqR0ldmezqVjaMpimiNRrz0hoi6IdW0W7kRJX6gvfp0tJh4"
    "keFOhVz1/3WVD6v3Q9fyyJ31fLvzawHvmhghhYZaT5bXH7oCudJeWiCYjI+omXUJhaLySAKNRa/mR/UWPCBMtM562bVnDFUOxE2NoHjPPaqPxPuGrJLmrkU5"
    "TgQY8uVZ0TgVghdxpqj8zkh2CEqzAnImpq3jqDc6DBxSEROCDO/p5DI+AvkPu7K1d1Aq/RA/8xFjx6dX58edvFL/9aajOnfhd0sVMxz0DIVeRBRuUrSRB6W/"
    "m3yLGIVMEtiAlyr/yS5iCZu+2maDga4Wb1j6JcsEE8ahPkaTHCLuYuXL3mbt125tEpv2DR7thJl46XFwGqimnwKb9wh10/9kpxeQyoBlHdbtppt59Tc0SiwG"
    "EwNvtoPCOJukb3Q94ONDN2bis3c6DSTVmzkc01aMSFYeu8g9dsc7zZL1ykyNIW0PR9J1OFkzyhNbeLod5ZPHsmf0FMg1d9qpNISB6B5EyI8hTxY7IyLrYEaV"
    "9z2+oxoHf73qk7qygzwq6FqP2zdaYTTTV+i0+rhkZRJIDOMnjOBKNdjlb/YcbHk/cS8ZK+62NAhomRj25G176svp18417/fjcm8tFXJHhLAinmv29GA3Bt1N"
    "Kr563Q3O45QkTGmVdr4XYxluusq6CJNM5GHFO4s5hdE5cIu29bETKHl92qRDlc75ajJ8qKRoP7BppvvBfaDJN4MFShgmIk+lYR/gIPRcuiKQUjpTWvFpef/P"
    "dKim/uW/zYj1e9B5+8CWmjn57F3ubBOkxFDluKCI0Rw/7DXZbanpDY7tHHL82lShaJQ7mhmyQj0glnmvlDIGq6IPhHxEfeN0mROj6rPkQOvYizqOuJ4O7fnx"
    "mG2moAcW2aHm4my4zydDVwVlUJjxnfIETdO9uGcM3VHD8jd33hrhn6DuLR/QvMxlvPy/YcTVGRarla00EGOlQq/DJpnixbs91GG/lsHc5OUiY7FGGz2EgYZ+"
    "urTpD8tCgclsUvOeybj+WW0omckko4gof/sNoyrj3nYo+WPAMmY7tGmxSbRBqUfHKii3nQlOq5GbypiDPLfL8YClcHfiAIrqop5G6dfsG4wQHMHLZYD9V/HX"
    "a2fxrOTqBS2/dYKdSwTF/o2QutYdTxT5/utj0WlgnXIgzwApVJaj9bgDvXPdiMom/ia55UPMBfTTrDgG+kvswPC4mK74nFgr0VDk+h5tVghV+pwvwEsfL/gj"
    "2svtITqvtFN5gBGpAYVajZDkkz7WxYvd4QhTYeNtKE1NEG406LhkMIlpcKwSR/m8V0gHe6qXH+NUWaDP5p39mGzl5meIjlUwBBE1+/Up9JLRPDys5EwYqtrm"
    "R02d1ryhTQEfYWCQUQRJeGr8P7p/KBaTY3DIjtXxwDqL+Q4MIMe4euZnL+kEdxd2xE0I4BkMjWB8zVpVQEhHbRTpmYCXn1MKhfheXgO0be9iWdfM5LlvGnWg"
    "5b8E1C4jnrCszuBCGOWTME/yX2RWNC3+tfbZ4P0PK/Wl/bqpd2qVU+WUbj7VRKWrAf/fFgwlxOUk/SOeR2OkUsmUqBjI94RcDhemP//ih5cRgGx/Z6NyZuYS"
    "/q/iD8afxg1eq0fTLwrtPUUhG9d+98UEiyN9LldpxrIcyz37y2wj0BuPBea/pqQM9mEhdb9BmPM8sO/gVG78jeP2h8cMJ+u4uDCOWPzkDhA/sqzyHq1MB0h+"
    "x01MUmBR5u0u0R1SO2s7PA/ObshWeqFxD6B3v/tYCc7thlKpLnMGzzUC3JKWO+bRmTD2O8Xs9pEZkDMI1LX1LkUskr827RJ0z/o0JA+VYdzvzYEwXi2uhaf3"
    "9Y+YTg2F5Y6k58RaiqwtkBiHYpp4NI8lskhZIPDPkMC8JA/uOchogfSgvZrvQ5w2Rwfji5//6mE0bHF+KKdAW28/LjIrFKrQkLRxhIyJZY4CGTlro9o2BU0u"
    "8Im51ejWE/GoUDsh+IWg3Tu3pETyeL6irIAnA1FosA5hLviStSezPvtZKlV4STyBW3vm2sOxK1zrZUjP3Hh/DKxHlvbS3D+c1IMHluYmth/etEk/sKizfgFk"
    "h8BJusdW9BKpxbGY8Op65nx/GIxM7dN4YzctH1c+IiCDq6HcmoLjP5MYpvBoWnc4EPT4sZ1s7Gn01zZypOBOLmaT7hHZ4r6HRN/iJlT9HvxbVU7s4HRfYmWd"
    "+A8SLJdrMd566zA1lpLBNI+goPckQoyk7S5XLc17eosHX4qG2/1k2jUsrS6eU5gVXkLp7Yz2Mc18jfV24kWJgt8MYk1j663i2zKOSMIq9U3sg5J/3qYsh3xY"
    "DrbOb5KVlPkYmX/kiItoMIJXZTiSXFks4+jo3I+l8vhrIlcFNFkRZm1pSt+JOLnYAUqpFSsMbXegc2gHMQ5/fXndke2UIf+WO0JOkc4K+QmaDXc2LfICvpVL"
    "k6QHLuTsfURPmJbyCbZn4glV8AMwjkWUMXYDW32Yn//eaxgzea9lZ4RRQXvQDXFvJL3HA2WGzgsfnESUo0+qaiPYyRHsUe43kFWMDhOidTitpElQMIbB3xMm"
    "FfqfGT8dXwzdSI0Q616UB9rs/lvmMLotsw+LpMipxUOMhaXbKW1Pqmzf3R3MkE5JiqHkl6Enlv5+XlATnoz7Kwwlg47C46uGVZRd18/7BJ6pa9ogMSYIHr7Q"
    "0xzsLkuYNHgpXsavrg4t+/BaBxcuc2Mmf4zryHldJCp3d+kphmds8eUbjkUeHa6sPFkvZ1vmhcroz0wX9Q73lHaSp6ntqBnCnnFIJ3QTTnYfJZh7lGeakhpI"
    "lzWEg23QyPjxSOOHnqK+S9H68W8JajT6fvArQIz3UetsyxuTr+KfAVV3b0nfJW00mab5eGezQhfhpWTzYrjXnXQS4/2Qni0HtRYsqOZb0SMsgEsh8NksoKrV"
    "Ly6oave//OfwfIclWTal6TJFAcrpX3T65z655l7TcBMbRAsQImcqwNArQORb4M4UeJrSpFh42xXvJ/DnGGgZOx7I0I6zzwbUvQGsX4PsRbcFxEfFj5zSI+4G"
    "Y07HwMTBLnP9xB5AjMXHAp0OyBlEmFYTD4hqP/BzVpmAtBHH90Uy4BUZFKw8v5JEc3If2h7TdUToc4R77wRRna+lmPC8/xxDSSFAACGWnZdwAj+wTQtXZK3T"
    "DpxQJJXGrUtwFNnFI9VHunZXHZ1Q85TQ2eRcNllHllIi/aREW9ZUwC3clxnamzfqY4cI0puiPmPVNfCWcDeAaFxCyj1z4dqrdNEVl6Fj6x2OW1EsGgwhWQNK"
    "Cj+XR3GGYLtG09f9cAtsK5R68sYSm6F9xwCXm+jmAf3GE68VZRu3Bs2IaokXYdfxqVOJiv9OTXaiGpel+4iYtd2+QAirKitzGVC+9U2YDWqzvjCbmnw8phcW"
    "lBaJovQYqFoJBKlVZziwcDLi6gzBKLp08kXfR7V6hyBNLQkLu7Uoiu0NEEhComgZK6xrnMYjObSxplNNREnYRwCsaXX37xSzSBWUgiM/u7ZEfgY8VoEr+JED"
    "QpIF9CbpfyqG6TYyxxOgeAZdAVM9jgidjAgRwgZDAgjO90s8VGapPJGsj36COpRHtRldZ21f0MfeJzCVxOVFNXSGzJCCWdfbuyDzxR/D5kix6TzgGTxKb20Q"
    "eKeYdlK1bQ7dAvXAiJIz9xGqyptK24HKztNPGpwPRkez2h9EukyLSHya8dAy+BTv+/Q8P+2mfDwv0TGdSV8710dmNbxG+Ft4A4aOkS5QcP1/5f/5DMVhPets"
    "QsioSt6Vt4yIkk4q7tE2FRX4blLo6hRAAV2kh+PTcATciu0CwTOdaND/6OOgeaYw7y12N2DZBRfTQonMoKW5KBbiBblZ5NQjoawj9PThteNxxvbsMvADH6Ye"
    "h/u/XVLRHLGRuf7stF9yI24GKBjO+GFw6RisA+ZgUuQ+xeKuEsExCPOSBEwL1JpN21c98nsZYacK2q9QokV6w73wgt7WUlznwrdQFeDEWZCnuK05x46ql7LX"
    "a5muJsCAYR1dyqHFJzh0/gSaSMzVvusBaPwjTL/Im3jKu/5rEVfuoCgtvQmFEfXefMHbpJJmpgqdph6dYAr1sYNVOjwq8sgbOLmYx6FK4bdejWvAeG89Yp0r"
    "vFvzQ7Th2wGFYavrGaNKVCoIP6AwzeBOx4xX/Dh0pO2cil1vt5G1uw5mUeywTEuoi9MzDmWZVX67ZIao1ff+Z4pe1P9wou4KEhTowElHOVx2Dt+KSJIjEMlB"
    "aBkbGngvFKPf0NyOUQxVGNjeq7yeXIV8E1bMFGt6yYVqZ+gV1ZmbrvPXpY1gbwM97klCi9EqVpFahters5Yrr/LKjP72kH0qc9m3vvtLkH2CL4Nq+geQYYzR"
    "BxXLar0y3+Djgegi1BUiYoiIYym72tT4ORTy++jdOJqiHW92osj+HqqCczzRuIVBnWXRShvQgO+iPvdWy1ghxdVNmbtbeb7mc9X/MkjSmh4RR/2dtidcJGFz"
    "8P3yRJV95ny2X/aBaCAINIcGa2ixCVwaPgHeUJJfUS6Oyy1FuaNAe5QfscJRUD18wy5pFWsqLfJOu3bMTNI0xksAYxIGgyA74cSXWHgh2GEPmCFovwCIuYLr"
    "6DVhZLebIzmEDmWKF2som8sDVQwp11LobzUu8kjiDej/KDglIEfvevbO2lqHE3Hi6M543v4LWe3nZ8N8kYhHIpOqh5C/uWtQ9LR0b32TZGUtPB6dvHTUGzIc"
    "6t38byklpNdXwmJpGYoUGMkg9imlSto6QahrIVH2snUE2vFdYlP3xtc7Kj+51uovnlpEjDnLKxJMsNVWo47R52ISrCsrVcjoj4fek/YvTPBK9hj/oOxdXOD6"
    "QFKDuMdqMpRw+Iu8enZ2bCZXeAIjMEiepCeQrJwfw/SWiHQp0VpMbTQeIMMbF17ahJe/N6NKWGGhszu9lGQRq9qvtn9Z2q+F7GAokRCIitdYCg94aobXPdVj"
    "QO3fhMDk2qOQ5zW7Y/iPPAz+R8LKL1MWkeqD1pa5G5TJa3ST0TrwsktkSgCUwzcLcbYPO74UAOVEurEdDSdtEk4PeS/r8/jsZ3cY5czMz21T5qoS4CAMp1cH"
    "JkHLuZUBx7+txBO67v1VJcLqucoTlmlDL2i3zH4/FnYYTIgzq/1BKcbEkv2ggPiSHL4xxF74dyehh2PWLhbt6G/MeGoNsT+UCAhNEznVvfPt1WPWFZIH05tY"
    "62dtTsWz3Af4z3yrtQTVSoihDj5LZVu/EKoR4pkBrLkd2xOfooNv9Z0DQ0VQXFbhUbMYwYpCzPrw/J4IwoejHPoR2HFH1I6jjydkeJhy+s+o1PJBDGc8Jhho"
    "kIKG6wm0/STbYpPHRO2l9mCLv/JA5VdKbdAySwf8EwdyCpb4aid04lLIX+Apie358ziev9N7kVe78PHHhut5Zw6OorZupJsFwQQc9dHdqypUgadsAo36XbOg"
    "iDsIhgjL4IkytB1p2nHz2RctowRkz634qyhisu+o2YN+fOAxN8ZR80+oPt7U2HfOww6WM2yWq5hEclYx4dkLGK//e+nR/GgcoHR11TI75k7LNlSguq9PYv/m"
    "r9yPVyX8tj3cmfow0HuQMfdgW20AG++yeZ6F8LL+eW519ZMYMSjVuDXtaqSpVO1JcUUSg78dN7xAZz5VvayYvxG0CzbvN5SoGeTlvpeQurn1lIk2/sorPbll"
    "3N9dklY1tPCaXaRMce+u/pu92xBmrASxYuwB10KUa2zVJZeZQTz0ynMEmfKf27C75DmAJf8Fq+Srcs9KqofhpENJ7MTo0tdJwwFQX2vCiIrIFShqXSSd92uj"
    "N58BfaHcFQyVTiAV/jRHP/7+S1CHZbBk70zYHQCVdfwmkvCV3CGIeIJQBMs31hXis1LQ+M8wt3Ux2oivkaPTIA4BxdCzDUPSrFwSk2tBwhwRWMhLbLvPrGBN"
    "zilA9mbkRiBjlT5DUsfj4sCOZ/4HXXjag06F71JiTcdqifmgDH3Gipfs4hm+vVkeI6v3qWUKyNCSCuvCgpZ4jhCKE6yJ3bLr9SpgWlCew8FLL/XF39K4x/qD"
    "hegiPxo4M5+VBeYoNj0/uWeV4VNh9tRg3Z9iLiCaqMtKb1FUFUM/kl1phZXjhrBK/59jpu6PVDxcQ3C4A2TRywQXZlqa87bJ6XjliLSkUM8sFf7L+O3NM4ZE"
    "891nF+HxcHs22MNzlPhBDZwcQd7JG6Ow28ekxtGU//6AKLigC2mwFXDcTGd/X8jXBg0mHC7HpR5iuYhcHS0VVZao3LDPYCSCtOFcjrA1ZBvi+8n+eQrBPzfy"
    "kQ/5t5CMGnIdE8qKdV8Q8EuEEcFWXrFldwne8ZhKhw6mYk0MLagB1YQw6Hp3BC8aHig/fN72PE7lHROm29hze8OZuThTs6qbVLJPYYoEMoGR6lJmwTmsEvI/"
    "gHaxXWox4rA0Ao5SHuRcW4j8G3qyuvNWLV0HzOMmJqMzXYHowAaSVvJ93pfAepwcz0yrKv9zN4h/MJ7OeX/+i7qDp/2sMU3okHbBuSBiRa++PtAQCjITfONT"
    "EUDHfrHZCjgaAywianAN98OMTGiW5x9Wdbg9dIgzFIVeWDXoVvfY/TxP5z9k2HcUp3K1JVhc2csNayKwnzfXf9jbvy27yX82cSQL8kZqFPT0JX3KDSd1Ocj1"
    "GH0x+YCp0i9GvbS6Lu6oWoRRJBMmnIQy5hl4R0f57zAkYCRvbNiDlKppPxvJH4JLHZpVN0Lwe/xhggSBdsMIFJQ2y+HKzsPlIFy00iS5O+4BJy5Lp+dChcZm"
    "2zUcXxoGLkRktk3Bl7D1sbWiLlSjhDshsfponT7qg1/WCI7Qsa7yhfzCard07ZNGyU02mVIzOmemwCgJwX4drIl3f3uskhgNhCpj32wM68L98eivbk92buIa"
    "XPJLuVyNQ9oXpyGKkQmciVbq+ijkdkScginMH/wfQNje0BdDZWCG6X22XNx7MN8XROBwrMl8o3yH5Yd0rjePGmtjC5gXmnlKWjIt9HLJGr3cI/f9Bt8siolU"
    "9qgx7EtfWWArYnbgxUZXVZQ6ynxLoYWduc6iknKVjRFKQyHKOHWRNDBz+NlDNwFcJc6TIiKxVBYQ5a0VHjuFscIhcChOkGRKbMP3oyuS6FBo2cSVV+qvQiGF"
    "eosmjWmCHiEL2rFf21UYj0+blRf0TGkViSezw93pEVMUgBHhvPt+/ROZh1KpQ8EQGD4BrJZ8QYGFK3FQNU39JMOnTx4zFDsJQ7tUfDiKNXXsmTpIKpIHM4wM"
    "zU0PyQJvvTsXtnMD9zu9mz0CETNIL7j7O/7gR+oIbEe+rpxRDhHmfojHUc02hINitD7fd7XUMAsHFYvy0XKMd4NREk80OgNGewXj87DQWisO7dLbUR0em74c"
    "5YjvKtPvEhPXwEiORXhk/2rPwa80IouP3rwzjtpg36gsSrxvalynrRBTg1srZHQkx1xsvko+MRRM71NMpm1ER/v2rxw5Hu69RdexU8sfuIfgO90f2fRwXxTL"
    "Q+HO3VGPkusp00EFHpINPbH/nQUtG7NSQMx/MmmNeu4atnRcCRWPfHgcCQetPh+XMeYUnJL31INf1mMQaaYPhWSChT3co3vcw4WGC/HzoLGRBs4By4vuhRRY"
    "HExfyDye+AX//boRqjRAXuchthO7rzuvNHJgj6XOeGdIhP9tEepNzYBuI2xx51+J+AogzenfZSdTHyBPW84mJIUs4L/R/Pz8cais2V0fvkkeECEjXeYRxfDJ"
    "YfHXgShZXrpzH8PSezzhuiwPIgfULvo46hJPMLq/g+ffF+PVaiF+EqfI6SVjnO+mwq6aRa+d8gF/wjGmpHzyV12zUwVutqU+fGQYAmiubROJveUpQ5O+SNuZ"
    "+z2CoaOZXEKYtSlG53IEurscCijDVx7/+bElR45th3QlcWRdpVQ/Sy9S8tBSwOeOqVO9xIE9zWDhgQ3NNdroxTBvd9iHoDS59vu6o6wJq4NwQEKsu2k2kewi"
    "ZYQB3pSuHwiFpI26bWDzxWsDWlJsgiLki97hiSGSPObstzrSy8Uw5kZoiAxN18E2QT2UQUauiRGA1DR9fzUw75XsqW9VRGZTFJAeYuJ1q1yHbK/kPKlqUi4A"
    "wTchDmAivZlHfkUQrM9PMbgZWJIOKBA00gqy4INF5Eip3FRcXHpz9BR7BOctjWO7bYJKfrPRtDXKG0w3qWApdH3O8w6fr7u5QW9Cu9gvQa3oXSywx54/N3yc"
    "AeD2DxLfH/2go0243iOpMzQBzBKmu/TlTyhKkHT9v+i1d4tnYzYVqDj6s7FuCQdPDE1NJKtsgPVuZzCyvMjaiFcoPnN6LzPnhAYQSGrg/oRWj+85RmzsA4h8"
    "J5sL479Pb/5WBh8wl+Ss2q14wf/fXGZ9rkNQIMakkhfxMbPFg7Xrj1xVdjTukcXDg+DMCNeFeuwrQnvI9Owdwn4wYpTykLL2NvTyhtsmWh8IhIA/ITT1yUsF"
    "m//Mpwr5+Gura0wcQA4HhmXXcTfhQKhUaxctdFyU/miQA7xFV5Ey9ewSzgHVCXT92DKHkcwJhX90qg+a82GXgXtTFPbMYGE1yNuqVS/c+qK0iwOlGGj+8h/2"
    "wkjngzRaiexweAv5oZylDWTCvWTXEhJODz7sj7L1rgNChdntgSCrP1LJI8YiAxZL5vAPRgsBaE91fSBoL9CjDX6EzLU2pCKBWjykAiecxEHYvLxg84tMJOqY"
    "TUGMH46XVCS/VT3MpCLkx8M5oVQy6jmNFV+IOxys6CdDskGff99Eto4RZ7D50hP6sE0l0eHoMMfc6d5kzkbULkT7D6hjWxmKPeL76w7TSFxKfom2hcbUovlf"
    "nMJg51RJ+HnrFrDHp150ttzazEYLan2Ls6mg0AShZphSlf/3fnzziYmdazeccaLjjMtuuMvmbUi+DYSn6TefMcCfKX4F2ASwmUxLlpPfOT2/aNkS2zWDnq/E"
    "D7cylafe5iHjmbYgekA6A+PNv8u8LLoXTTnrY85TJl2x+Gd/DSvjQ2cFHKZl/PrxUEKmhLU0BfpCSNZOCmQz9kZOFTm1sVdtzvwqC3vEuOWB8WcHgWRhCPfr"
    "nHy4b33KILWlDViIelgPrRKSLXQ2+buMlKSWh64EhC4hfvQhle0hzCbjGUHpdqKbx++Lud1jZtmMR84vveTr5PWKaxNG93edLUDiFb292E7iBbn95uj+cQFs"
    "teAMMs238gtFdiXN+WnAElr0RqSuZ7AVP0xVrQvNHHESNvVSB2O20qZ7UIr2/VDGASFP4JqlgAwsARhEwbYIC1E3wqGTlBs6z/Q2SQHxq3GOU4psxg3kmJzy"
    "ME96I31yc1JXlG6IPoX7HlgJizgIRNd6syA4MEi36nEqEccMmFIh3wfPZsg4joYPEhFqR4GsgR69lU7eGD0lqo7x7KUxVx/YH7W0SDr6eQgbDM2aBhfBPnSV"
    "KKYw0TMQzX+USHMbv7ov8kf2mhqehBmY6rP5OMLNET3CmQRiJi+OMQDFOM0kYDIrDGmuR5MkVsJbG2hJpEoCgyRyfPdS/EcqbDQoJgLaXGvsxnUDZl+Cl8mO"
    "cDJ6yHiwOXbmyR9V/2Huzm4WzB0KJUA8y1XWGRVcpkrDutVuuebj426CIpko+YLAxwKAyZ53wYo8jDvMgSeTZePcF1ZXnGTB9eWvHwaVPIRUWCS3jHIXToIn"
    "H88ycqcbRl3f6SaCFSyhrwfnYWo2+uDTE56jjoyhZ4iLAnDykPSSUKEaH88yTPxo58vo2ECkZppFkcqX9zhar+bQG8Iv+E6rKfVNr4wDpTZKkcJffVHnNqY/"
    "Zzv8MG/ByvsjiEi3x9p0gGWixeIk7Cmo4Vg1gnN6J3jq+uWAuk4/rFb+a5uJNORzA1xPOC2d64HJT/tIH1axA3zEAKhSNa9bLnjAOxTXyE44SIL4wUcSvZey"
    "TX1pB8tbYplLg5Mn/4F1jyHXhBfmThjda2nkI8f3H15tutSq/XtrowXV32De17TtjHi6T8VcBypGzEM3CqOenbhkm4zgw8wROhFUvnvQbxuooXGAh3Sph4Z/"
    "SGjv+oexuIIP6x93Om+AoBBwLp4wD6dp5CQx5WwANRXYN8g0Qhj6NlEPa4ZlrQk+T6DwewYujktG6d32GzlHgCcJ9nXDL6Wv9bIzCS98Qh1JeVUCUh11TNEv"
    "K0uWQ1f6ZSNZSf4bcR1sRz/acM3WnjWm0mpvHF0/uetsZph6J3nmABhQ+fVBAKbnsxLUAYgncWedWXbF04FWCUEnZIYCcW4tg0LRK4NEDiTNGcmaBrfbG+Rz"
    "5WtjiFODn2gvG1Xzr4ePM9KTShalu375E0s0s8UtZouznugHgPYWrd80X2F9yDZJPLclhnaxFtTmmmlFHjgEWJEtFqjvHJ/1x9sUBORkrslvJXwOqkxJ4ecP"
    "UMK63aifomVCsfH5SxPoH0vNM+3LsMUHAPnT1L9HY4YLFRkGLGYLQ+8euc96lEkT9tXyZZlyHuxYJr+8U9dHN1KM0l1eAAAA2Y9yyImfIswAAZ6iAe/ABLeM"
    "SzixxGf7AgAAAAAEWVo="
)

DEFAULT_FILE = os.environ.get("FARMACOS_XLSX", "Reporte_Farmacos__NUEVA_.xlsx")

MES_NUM = {"ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6,
           "JUL": 7, "AGO": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DIC": 12}
MES_ABR = {v: k for k, v in MES_NUM.items()}
MES_NOMBRE_MAP = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9, "OCTUBRE": 10,
    "NOVIEMBRE": 11, "DICIEMBRE": 12,
    "ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6,
    "JUL": 7, "AGO": 8, "SEP": 9, "SEPT": 9, "OCT": 10, "NOV": 11, "DIC": 12,
}

BODEGAS_FARM_STOCK = {
    "BODEGA FARMACOS 2 HV": "FARMACOS 2",
    "BODEGA FARMACOS HV": "FARMACOS",
    "BODEGA ACTIVA FARMACIA AT. ABIERTA HV": "STOCK AT ABIERTA",
    "BODEGA ACTIVA FARMACIA AT. CERRADA HV": "STOCK AT CERRADA",
    "BODEGA STOCKLIGTH AT. ABIERTA": "STOCKLIGTH ABIERTA",
    "BODEGA STOCKLIGTH AT. CERRADA": "STOCKTIGTH CERADA",
    "BODEGA BOTIQUIN FARMACIA URGENCIA HV": "BOTIQUIN URGENCIA",
    "BODEGA ACTIVA FARMACIA C.E.M.": "BODEGA CEM",
}

ESTADOS = ["QS", "CR", "OK", "SS"]
ESTADO_TXT = {"QS": "QS (quiebre de stock)", "CR": "CR (crítico)",
              "OK": "OK (normal)", "SS": "SS (sobre stock)"}
ESTADO_ICON = {"QS": "🔴 QS", "CR": "🟠 CR", "OK": "🟢 OK", "SS": "🔵 SS"}
ESTADO_COLOR = {"QS": "#B3261E", "CR": "#E8890C", "OK": "#2E7D5B", "SS": "#3B6EA8"}
ESTADO_TXT.update({"SM": "SM (con stock, sin salidas)", "IN": "IN (inactivo)"})
ESTADO_ICON.update({"SM": "⚪ SM", "IN": "⚫ IN"})
ESTADO_COLOR.update({"SM": "#8A8A8A", "IN": "#C9C9C9"})
ROT_ORDEN = ["SR", "BR", "MIR", "MEDR", "ALR", "GR", "MAXR", "OK"]
CEN_RIESGO = ["FALTANTE", "SUSPENDIDO", "ELIMINADO"]
Z80 = 1.2816
MIN_TRAIN = 3
HORIZONTE_FC = 18

st.markdown(
    """
<style>
.kpi{border-radius:8px;padding:.8rem 1rem;background:rgba(128,128,128,.10);margin-bottom:.4rem}
.kpi-l{font-size:.85rem;opacity:.75}
.kpi-v{font-size:2rem;font-weight:700;line-height:1.15}
.kpi-s{font-size:.8rem;opacity:.65}
</style>
""",
    unsafe_allow_html=True,
)


def n0(x):
    if x is None or (isinstance(x, float) and (np.isnan(x) or np.isinf(x))):
        return "–"
    return f"{x:,.0f}".replace(",", ".")


def clp(x):
    s = n0(x)
    return s if s == "–" else "$ " + s


def kpi(label, value, sub="", color="#888888"):
    st.markdown(
        f'<div class="kpi" style="border-left:6px solid {color}">'
        f'<div class="kpi-l">{label}</div><div class="kpi-v">{value}</div>'
        f'<div class="kpi-s">{sub}</div></div>',
        unsafe_allow_html=True,
    )


def parse_mes(label):
    m = re.match(r"^(?:PROG )?([A-Z]{3})-(\d{2})(?: CONS)?$", str(label).strip())
    if m and m.group(1) in MES_NUM:
        return (2000 + int(m.group(2)), MES_NUM[m.group(1)])
    return None


def mes_label(ym):
    return f"{MES_ABR[ym[1]]}-{str(ym[0])[2:]}"


def next_month(ym):
    return (ym[0] + (1 if ym[1] == 12 else 0), ym[1] % 12 + 1)


def mes_ts(ym):
    return pd.Timestamp(ym[0], ym[1], 1)


def parse_clp(v):
    if isinstance(v, (int, float, np.number)):
        return float(v)
    s = re.sub(r"[^\d,\-]", "", str(v)).replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return np.nan


def rot_from_consumo(x):
    return np.select(
        [x <= 0, x < 10, x < 100, x < 1000, x < 5000, x < 10000, x < 190000],
        ["SR", "BR", "MIR", "MEDR", "ALR", "GR", "MAXR"],
        "OK",
    )


def _fecha_dmy(d, m, y):
    try:
        return date(int(y), int(m), int(d))
    except (ValueError, TypeError):
        return None


def leer_con_encabezado(file_bytes, clave="Código"):
    raw = pd.read_excel(io.BytesIO(file_bytes), header=None)
    hdr = None
    for i in range(min(25, len(raw))):
        if str(raw.iloc[i, 0]).strip().lower() == clave.lower():
            hdr = i
            break
    if hdr is None:
        raise ValueError(f"No encontré la fila de títulos que empieza con '{clave}'.")
    d = raw.iloc[hdr + 1:].copy()
    d.columns = [str(x).strip() for x in raw.iloc[hdr].tolist()]
    col_real = d.columns[0]
    d = d.dropna(subset=[col_real]).reset_index(drop=True)
    texto = " ".join(str(x) for x in raw.iloc[:hdr, 0].dropna().tolist())
    return texto, d


@st.cache_data(show_spinner=False)
def clasificar(nombre: str, file_bytes: bytes):
    nom_up = nombre.upper()
    if nom_up.endswith(".CSV"):
        try:
            inicio = file_bytes[:1500].decode("latin1", errors="ignore").upper()
            if "CENABAST" in inicio or "SEMAFORO;" in inicio or "INTERMEDIACION" in inicio:
                return "icp_cenabast"
        except Exception:
            pass
        return None

    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        if "ARSENAL_FARMACIA" in wb.sheetnames:
            return "farmacos"
        ws = wb[wb.sheetnames[0]]
        filas = []
        for fila in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            filas.append([str(c or "").strip() for c in fila[:6]])
        a1 = (filas[0][0] if filas and filas[0] else "").lower()
        a2 = (filas[1][0] if len(filas) > 1 and filas[1] else "").lower()

        if a1.startswith("reporte de licitaciones"):
            return "licitaciones_raw"
        if a1.startswith("reporte de productos sin recepcionar"):
            return "sin_recepcionar"
        if a1.startswith("reporte de stock"):
            if "todas las bodegas" in a2:
                return "stock_centralizado_raw"
            return "vencimientos"
        if a1.startswith("stock"):
            return "stock_insumos"
        if a1.startswith("movimientos"):
            es_farm = any(k in a1 for k in ("farmacia", "stockligth", "stocklight", "botiquin", "c.e.m"))
            return "consumo_farm_raw" if es_farm else "salidas_insumos"
        if "programacion" in nom_up or any("cant. programada" in " ".join(f).lower() for f in filas):
            return "programacion_raw"
    except Exception:
        pass
    return None


def reunir_archivos(ups):
    def vacio():
        return {
            "farmacos": None,
            "stock_insumos": None,
            "vencimientos": None,
            "salidas_insumos": [],
            "stock_centralizado_raw": None,
            "licitaciones_raw": None,
            "icp_cenabast": [],
            "consumo_farm_raw": [],
            "programacion_raw": [],
        }

    def registrar(res, vistos, nombre, b):
        if nombre.upper().endswith(".ZIP"):
            try:
                with zipfile.ZipFile(io.BytesIO(b), "r") as z:
                    for zinfo in z.infolist():
                        if not zinfo.is_dir() and not os.path.basename(zinfo.filename).startswith("."):
                            registrar(res, vistos, zinfo.filename, z.read(zinfo.filename))
            except Exception:
                pass
            return
        h = hash(b)
        if h in vistos:
            return
        vistos.add(h)
        t = clasificar(nombre, b)
        if t in ("salidas_insumos", "icp_cenabast", "consumo_farm_raw", "programacion_raw"):
            res[t].append((nombre, b))
        elif t in ("farmacos", "stock_insumos", "vencimientos", "stock_centralizado_raw", "licitaciones_raw"):
            if res[t] is None:
                res[t] = (nombre, b)

    subidos, locales = vacio(), vacio()
    vistos_s, vistos_l = set(), set()
    for u in ups or []:
        registrar(subidos, vistos_s, u.name, u.getvalue())

    rutas = []
    if os.path.exists(DEFAULT_FILE):
        rutas.append(DEFAULT_FILE)
    rutas += sorted(glob.glob("*.xlsx") + glob.glob("*.csv") + glob.glob("*.zip"), key=os.path.getmtime, reverse=True)
    for r in rutas:
        try:
            with open(r, "rb") as fh:
                registrar(locales, vistos_l, os.path.basename(r), fh.read())
        except OSError:
            pass

    res = vacio()
    for t in ("farmacos", "stock_insumos", "vencimientos", "stock_centralizado_raw", "licitaciones_raw"):
        res[t] = subidos[t] or locales[t]
    for t in ("salidas_insumos", "icp_cenabast", "consumo_farm_raw", "programacion_raw"):
        res[t] = subidos[t] or locales[t]
    return res


def hay_archivos_crudos_farmacos(archivos: dict) -> bool:
    return any([
        archivos["stock_centralizado_raw"] is not None,
        archivos["licitaciones_raw"] is not None,
        bool(archivos["icp_cenabast"]),
        bool(archivos["consumo_farm_raw"]),
        bool(archivos["programacion_raw"]),
    ])


@st.cache_data(show_spinner=False)
def cargar_base_semilla():
    raw_csv = lzma.decompress(base64.b64decode(ARSENAL_SEED_B64))
    ars = pd.read_csv(io.BytesIO(raw_csv))
    ars["dato"] = ars["dato"].astype(str).str.strip()
    texto = {"dato", "DESCRIPCION", "SEGUIMIENTO"}
    for c in ars.columns:
        if c not in texto:
            ars[c] = pd.to_numeric(ars[c], errors="coerce")
    for col_st in list(BODEGAS_FARM_STOCK.values()) + ["TOTAL BODEGAS", "TOTAL FARMACIAS", "STOCK TOTAL", "STOCK SIN STOCKLIGHT"]:
        ars[col_st] = 0.0
    conv = pd.DataFrame(columns=["dato", "ID LICITACION", "SALDO", "VCTO"])
    venc = pd.DataFrame(columns=["Bodega", "Código", "Producto", "Lote", "Fecha Vencimiento", "Cantidad", "Precio Lote", "Total"])
    lic = pd.DataFrame(columns=["Código", "Licitación", "Saldo", "Proveedor", "Vencimiento"])
    precios = pd.Series(dtype=float)
    return {"ars": ars, "conv": conv, "venc": venc, "precios": precios, "fecha_txt": None, "licit": lic}


@st.cache_data(show_spinner="Leyendo el Excel de fármacos…")
def load_data(file_bytes: bytes):
    xls = pd.ExcelFile(io.BytesIO(file_bytes))
    ars = xls.parse("ARSENAL_FARMACIA")
    ars.columns = [str(c).strip() for c in ars.columns]
    ars = ars[ars["dato"].notna()].copy()
    ars["dato"] = ars["dato"].astype(str).str.strip()
    ars = ars.drop_duplicates("dato").reset_index(drop=True)
    texto = {"dato", "DESCRIPCION", "ESTADO", "ROTACION", "ESTADO CENABAST",
             "FECHA ENTREGA", "ENTREGA EN DIA", "SEGUIMIENTO", "CTRL2"}
    for c in ars.columns:
        if c not in texto:
            ars[c] = pd.to_numeric(ars[c], errors="coerce")

    try:
        conv = xls.parse("CONVENIOS_VIGENTES")
        conv.columns = [str(c).strip() for c in conv.columns]
        conv["dato"] = conv["dato"].astype(str).str.strip()
        conv = conv.drop_duplicates("dato")
        conv["VCTO"] = pd.to_datetime(conv["VCTO"], errors="coerce")
        conv["SALDO"] = pd.to_numeric(conv["SALDO"], errors="coerce")
        conv = conv[["dato", "ID LICITACION", "SALDO", "VCTO"]]
    except Exception:
        conv = pd.DataFrame(columns=["dato", "ID LICITACION", "SALDO", "VCTO"])

    try:
        venc = xls.parse("VENCIMIENTOS", header=2)
        venc.columns = [str(c).strip() for c in venc.columns]
        venc["Fecha Vencimiento"] = pd.to_datetime(venc["Fecha Vencimiento"], errors="coerce")
        venc = venc.dropna(subset=["Fecha Vencimiento"]).copy()
        venc["Código"] = venc["Código"].astype(str).str.strip()
        venc["Lote"] = venc["Lote"].astype(str)
        venc["Cantidad"] = pd.to_numeric(venc["Cantidad"], errors="coerce").fillna(0)
        venc["Precio Lote"] = venc["Precio Lote"].map(parse_clp)
        venc["Total"] = venc["Total"].map(parse_clp)
    except Exception:
        venc = pd.DataFrame(columns=["Bodega", "Código", "Producto", "Lote",
                                     "Fecha Vencimiento", "Cantidad", "Precio Lote", "Total"])

    try:
        lic = xls.parse("LICITACIONES")
        lic.columns = [str(c).strip() for c in lic.columns]
        lic["Código"] = lic["Código"].astype(str).str.strip()
        lic["Saldo"] = pd.to_numeric(lic["Saldo"], errors="coerce")
        lic["Vencimiento"] = pd.to_datetime(lic["Vencimiento"], errors="coerce")
        lic = lic[["Código", "Licitación", "Saldo", "Proveedor", "Vencimiento"]]
    except Exception:
        lic = pd.DataFrame(columns=["Código", "Licitación", "Saldo", "Proveedor", "Vencimiento"])

    precios = pd.Series(dtype=float)
    try:
        sc = xls.parse("STOCK CENTRALIZADO", usecols=["Código", "Cantidad", "Precio"])
        sc["Código"] = sc["Código"].astype(str).str.strip()
        sc = sc[sc["Código"].isin(ars["dato"])].copy()
        sc["Cantidad"] = pd.to_numeric(sc["Cantidad"], errors="coerce")
        sc["Precio"] = pd.to_numeric(sc["Precio"], errors="coerce")
        pos = sc[(sc["Cantidad"] > 0) & (sc["Precio"] > 0)].copy()
        pos["pq"] = pos["Precio"] * sc["Cantidad"]
        g = pos.groupby("Código")[["pq", "Cantidad"]].sum()
        p1 = g["pq"] / g["Cantidad"]
        p2 = sc[sc["Precio"] > 0].groupby("Código")["Precio"].mean()
        precios = p1.combine_first(p2)
    except Exception:
        pass

    fecha_txt = None
    try:
        rep = xls.parse("REPORTE", header=None)
        for v in rep.to_numpy().ravel():
            m = re.search(r"ACTUALIZACI[OÓ]N\s*(\d{1,2})/(\d{1,2})", str(v), flags=re.I)
            if m:
                fecha_txt = (int(m.group(1)), int(m.group(2)))
                break
    except Exception:
        pass

    return {"ars": ars, "conv": conv, "venc": venc, "precios": precios, "fecha_txt": fecha_txt, "licit": lic}


def _extraer_ym_nombre(nombre, texto="", anio_defecto=None):
    s = f"{nombre} {texto}".upper()
    m = re.search(r"desde\s*(\d{1,2})/(\d{1,2})/(\d{4})", s, flags=re.I)
    if m:
        return (int(m.group(3)), int(m.group(2)))
    for mes_nom, m_num in sorted(MES_NOMBRE_MAP.items(), key=lambda x: -len(x[0])):
        pat = rf"(?:^|[^A-Z]){mes_nom}(?:[^A-Z]|$)"
        if re.search(pat, s):
            ym = re.search(r"20(2\d|3\d)", s)
            if ym:
                yr = 2000 + int(ym.group(1))
            else:
                yr = anio_defecto or date.today().year
            return (yr, m_num)
    return None


def cruzar_datos_crudos(data: dict, archivos: dict):
    ars = data["ars"].copy()
    conv = data["conv"].copy()
    lic = data["licit"].copy()
    venc = data["venc"].copy()
    precios = data["precios"].copy()
    fecha_txt = data["fecha_txt"]
    log_cruces = []

    if archivos["stock_centralizado_raw"] is not None:
        nom_s, b_s = archivos["stock_centralizado_raw"]
        try:
            texto_s, sc = leer_con_encabezado(b_s, clave="Código")
            sc["Código"] = sc["Código"].astype(str).str.strip()
            sc["Bodega_norm"] = sc["Bodega"].astype(str).str.replace(r"\s+", " ", regex=True).str.strip().str.upper()
            sc["Cantidad"] = pd.to_numeric(sc["Cantidad"], errors="coerce").fillna(0)
            sc["Precio"] = pd.to_numeric(sc["Precio"], errors="coerce").fillna(0)

            piv_s = (
                sc[sc["Bodega_norm"].isin(BODEGAS_FARM_STOCK)]
                .groupby(["Código", "Bodega_norm"])["Cantidad"]
                .sum()
                .unstack(fill_value=0)
                .rename(columns=BODEGAS_FARM_STOCK)
            )
            for col_dest in BODEGAS_FARM_STOCK.values():
                if col_dest in piv_s.columns:
                    ars[col_dest] = ars["dato"].map(piv_s[col_dest]).fillna(0)
                else:
                    ars[col_dest] = 0.0

            ars["TOTAL BODEGAS"] = ars["FARMACOS"].fillna(0) + ars["FARMACOS 2"].fillna(0)
            cols_farm = ["STOCK AT ABIERTA", "STOCK AT CERRADA", "STOCKLIGTH ABIERTA",
                         "STOCKTIGTH CERADA", "BOTIQUIN URGENCIA", "BODEGA CEM"]
            ars["TOTAL FARMACIAS"] = ars[cols_farm].fillna(0).sum(axis=1)
            ars["STOCK TOTAL"] = ars["TOTAL BODEGAS"] + ars["TOTAL FARMACIAS"]
            ars["STOCK SIN STOCKLIGHT"] = (
                ars["STOCK TOTAL"] - ars["STOCKLIGTH ABIERTA"].fillna(0) - ars["STOCKTIGTH CERADA"].fillna(0)
            )

            sc_f = sc[sc["Código"].isin(ars["dato"])].copy()
            pos = sc_f[(sc_f["Cantidad"] > 0) & (sc_f["Precio"] > 0)].copy()
            pos["pq"] = pos["Precio"] * pos["Cantidad"]
            g = pos.groupby("Código")[["pq", "Cantidad"]].sum()
            p1 = g["pq"] / g["Cantidad"]
            p2 = sc_f[sc["Precio"] > 0].groupby("Código")["Precio"].mean() if len(sc_f) else pd.Series(dtype=float)
            precios = p1.combine_first(p2)

            m_f = re.search(r"Obtenido\s*:\s*(\d{1,2})/(\d{1,2})/(\d{4})", texto_s, flags=re.I)
            if m_f:
                fecha_txt = (int(m_f.group(1)), int(m_f.group(2)))
            log_cruces.append(f"Stock centralizado ({os.path.basename(nom_s)})")
        except Exception as e:
            st.sidebar.warning(f"No se pudo cruzar Stock Centralizado: {e}")

    if archivos["licitaciones_raw"] is not None:
        nom_l, b_l = archivos["licitaciones_raw"]
        try:
            _, df_l = leer_con_encabezado(b_l, clave="Licitación")
            df_l["Código"] = df_l["Código"].astype(str).str.strip()
            df_l["Saldo"] = pd.to_numeric(df_l["Saldo"], errors="coerce")
            vcto = pd.to_datetime(df_l["Vencimiento"], format="%d/%m/%Y", errors="coerce")
            if vcto.isna().any():
                vcto[vcto.isna()] = pd.to_datetime(df_l.loc[vcto.isna(), "Vencimiento"], errors="coerce")
            df_l["Vencimiento"] = vcto
            lic = df_l[["Código", "Licitación", "Saldo", "Proveedor", "Vencimiento"]].copy()

            prim_lic = lic.drop_duplicates("Código", keep="first").set_index("Código")
            conv = pd.DataFrame({
                "dato": ars["dato"],
                "ID LICITACION": ars["dato"].map(prim_lic["Licitación"]),
                "SALDO": pd.to_numeric(ars["dato"].map(prim_lic["Saldo"]), errors="coerce"),
                "VCTO": pd.to_datetime(ars["dato"].map(prim_lic["Vencimiento"]), errors="coerce"),
            })
            log_cruces.append(f"Licitaciones y Convenios ({os.path.basename(nom_l)})")
        except Exception as e:
            st.sidebar.warning(f"No se pudo cruzar Licitaciones: {e}")

    if archivos["icp_cenabast"]:
        def _orden_icp(item):
            ym = _extraer_ym_nombre(item[0])
            return ym if ym else (0, 0)

        nom_icp, b_icp = sorted(archivos["icp_cenabast"], key=_orden_icp)[-1]
        try:
            txt_icp = b_icp.decode("latin1", errors="ignore")
            lineas = txt_icp.splitlines()
            hdr_idx = next(i for i, l in enumerate(lineas[:15]) if "SEMAFORO" in l.upper() and ";" in l)
            df_icp = pd.read_csv(io.StringIO("\n".join(lineas[hdr_idx:])), sep=";")
            df_icp.columns = [str(c).strip().upper() for c in df_icp.columns]
            df_icp["CODIGO GENERICO"] = pd.to_numeric(df_icp["CODIGO GENERICO"], errors="coerce")
            df_icp = df_icp.dropna(subset=["CODIGO GENERICO"]).drop_duplicates("CODIGO GENERICO", keep="first")
            icp_map = df_icp.set_index("CODIGO GENERICO")

            zgen_s = pd.to_numeric(ars["ZGEN"], errors="coerce")
            if "CANTIDAD UNITARIA PROGRAMADA" in icp_map.columns:
                ars["CANT PROGRAMADA CENABAST"] = pd.to_numeric(
                    zgen_s.map(icp_map["CANTIDAD UNITARIA PROGRAMADA"]), errors="coerce"
                )
            if "Q FALTANTE" in icp_map.columns:
                ars["CANT FALTANTE CENABAST"] = pd.to_numeric(
                    zgen_s.map(icp_map["Q FALTANTE"]), errors="coerce"
                )
            if "CANTIDAD UNITARIA A DESPACHAR" in icp_map.columns:
                ars["CANT ENTREGA CENABAST"] = pd.to_numeric(
                    zgen_s.map(icp_map["CANTIDAD UNITARIA A DESPACHAR"]), errors="coerce"
                )
            if "ESTADO DEL MATERIAL" in icp_map.columns:
                ars["ESTADO CENABAST"] = zgen_s.map(icp_map["ESTADO DEL MATERIAL"].astype(str).str.strip())
            log_cruces.append(f"ICP Cenabast ({os.path.basename(nom_icp)})")
        except Exception as e:
            st.sidebar.warning(f"No se pudo cruzar ICP Cenabast: {e}")

    if archivos["consumo_farm_raw"]:
        por_mes = {}
        n_archivos_mes = {}
        for nom_c, b_c in archivos["consumo_farm_raw"]:
            try:
                texto_c, df_c = leer_con_encabezado(b_c, clave="Código")
                ym = _extraer_ym_nombre(nom_c, texto_c)
                if ym is None:
                    continue
                n_archivos_mes[ym] = n_archivos_mes.get(ym, 0) + 1
                df_c["Código"] = df_c["Código"].astype(str).str.strip()
                c_col = "cantidad" if "cantidad" in df_c.columns else "Cantidad"
                df_c["cant"] = pd.to_numeric(df_c[c_col], errors="coerce").fillna(0)
                if "Tipo Documento" in df_c.columns:
                    td = df_c["Tipo Documento"].astype(str).str.strip().str.upper()
                    dest = (
                        df_c["Bodega Destino"].astype(str).str.strip().str.upper()
                        if "Bodega Destino" in df_c.columns
                        else pd.Series("", index=df_c.index)
                    )
                    cond = (td != "SOLICITUD AJUSTE") & ~((td == "TRASPASO") & dest.str.startswith("VILLARRICA HOSP."))
                    df_c = df_c[cond]
                s_mes = df_c.groupby("Código")["cant"].sum()
                por_mes[ym] = por_mes.get(ym, pd.Series(0.0, index=ars["dato"])).add(s_mes, fill_value=0)
            except Exception:
                pass
        meses_act = []
        for ym, serie_m in sorted(por_mes.items()):
            col_c = f"{mes_label(ym)} CONS"
            if col_c in ars.columns and n_archivos_mes.get(ym, 0) < 6:
                ars[col_c] = np.maximum(ars[col_c].fillna(0), ars["dato"].map(serie_m).fillna(0))
            else:
                ars[col_c] = ars["dato"].map(serie_m).fillna(0)
            meses_act.append(mes_label(ym))
        if meses_act:
            log_cruces.append(f"Consumos de farmacia ({', '.join(meses_act)})")

    if archivos["programacion_raw"]:
        prog_mes = {}
        for nom_p, b_p in archivos["programacion_raw"]:
            try:
                ym = _extraer_ym_nombre(nom_p)
                if ym is None:
                    continue
                df_p = pd.read_excel(io.BytesIO(b_p))
                df_p.columns = [str(c).strip() for c in df_p.columns]
                c_cod = next((c for c in df_p.columns if "cód" in c.lower() or "cod" in c.lower()), df_p.columns[0])
                c_can = next((c for c in df_p.columns if "cant" in c.lower() or "prog" in c.lower()), df_p.columns[-1])
                df_p[c_cod] = df_p[c_cod].astype(str).str.strip()
                df_p[c_can] = pd.to_numeric(df_p[c_can], errors="coerce").fillna(0)
                s_p = df_p.groupby(c_cod)[c_can].sum()
                prog_mes[ym] = prog_mes.get(ym, pd.Series(0.0, index=ars["dato"])).add(s_p, fill_value=0)
            except Exception:
                pass
        for ym, serie_p in sorted(prog_mes.items()):
            col_p = f"PROG {mes_label(ym)}"
            ars[col_p] = ars["dato"].map(serie_p).fillna(0)
        if prog_mes:
            log_cruces.append(f"Programación ({', '.join(mes_label(k) for k in sorted(prog_mes))})")

    if venc.empty and archivos["vencimientos"] is not None:
        try:
            _, d_v = leer_con_encabezado(archivos["vencimientos"][1], clave="Establecimiento")
            d_v["Código"] = d_v["Código"].astype(str).str.strip()
            d_v = d_v[d_v["Código"].isin(set(ars["dato"]))].copy()
            fv = pd.to_datetime(d_v["Fecha Vencimiento"], format="%d/%m/%Y", errors="coerce")
            if fv.isna().any():
                fv[fv.isna()] = pd.to_datetime(d_v.loc[fv.isna(), "Fecha Vencimiento"], errors="coerce")
            d_v["Fecha Vencimiento"] = fv
            d_v = d_v.dropna(subset=["Fecha Vencimiento"]).copy()
            d_v["Cantidad"] = pd.to_numeric(d_v["Cantidad"], errors="coerce").fillna(0)
            d_v["Precio Lote"] = d_v["Precio Lote"].map(parse_clp)
            d_v["Total"] = d_v["Total"].map(parse_clp)
            venc = d_v[["Bodega", "Código", "Producto", "Lote", "Fecha Vencimiento", "Cantidad", "Precio Lote", "Total"]].reset_index(drop=True)
        except Exception:
            pass

    if log_cruces:
        cons_cols = sorted(
            [c for c in ars.columns if re.match(r"^[A-Z]{3}-\d{2} CONS$", c) and parse_mes(c)],
            key=lambda c: parse_mes(c),
        )
        ult3 = cons_cols[-3:] if len(cons_cols) >= 3 else cons_cols
        mat3 = ars[ult3].apply(pd.to_numeric, errors="coerce").fillna(0).to_numpy(float)
        con_nz = np.where(mat3 != 0, mat3, np.nan)
        with np.errstate(invalid="ignore"):
            prom = np.nan_to_num(np.nanmean(con_nz, axis=1), nan=0.0)
        ars["CONSUMO PROMEDIO"] = prom
        if mat3.shape[1] == 3:
            mayor_c = np.maximum(mat3[:, 0], mat3[:, 2])
        else:
            mayor_c = mat3.max(axis=1)
        ars["MAYOR CONSUMO"] = mayor_c
        ars["CONSUMO DIARIO"] = ars["CONSUMO PROMEDIO"] / 30.0
        ars["CONSUMO SEMANAL"] = ars["CONSUMO DIARIO"] * 7.0
        ars["CONSUMO QUINCENAL"] = ars["CONSUMO DIARIO"] * 15.0
        ars["STOCK CRITICO"] = ars["CONSUMO DIARIO"] * 20.0
        ars["STOCK IDEAL"] = ars["CONSUMO DIARIO"] * 60.0

        cd = ars["CONSUMO DIARIO"].replace(0, np.nan)
        ars["AUTONOMIA DIAS (BODEGAS)"] = (ars["TOTAL BODEGAS"] / cd).fillna(0)
        ars["AUTONOMIA DIAS (HOSPITAL)"] = (ars["STOCK TOTAL"] / cd).fillna(0)

        q_sin = ars["STOCK SIN STOCKLIGHT"].fillna(0).to_numpy(float)
        ad_prom = ars["CONSUMO PROMEDIO"].fillna(0).to_numpy(float)
        ar_prog = (
            pd.to_numeric(ars["PROG JUL-26"], errors="coerce").fillna(0).to_numpy(float)
            if "PROG JUL-26" in ars.columns
            else np.zeros(len(ars))
        )
        aut_m = np.where(
            ad_prom > 0,
            q_sin / np.where(ad_prom > 0, ad_prom, 1.0),
            np.where(ar_prog > 0, q_sin / np.where(ar_prog > 0, ar_prog, 1.0), np.where(q_sin == 0, 0.0, 2.0)),
        )
        ars["AUTONOMIA MENSUAL (HOSP - STOCKLIGHT)"] = aut_m
        ars["AUTONOMIA DIAS (HOSP - STOCKLIGHT)"] = aut_m * 30.0
        ars["ESTADO"] = np.select(
            [ars["AUTONOMIA DIAS (HOSP - STOCKLIGHT)"] <= 0,
             ars["AUTONOMIA DIAS (HOSP - STOCKLIGHT)"] < 30,
             ars["AUTONOMIA DIAS (HOSP - STOCKLIGHT)"] < 360],
            ["QS", "CR", "OK"],
            "SS",
        )
        ars["ROTACION"] = rot_from_consumo(ars["MAYOR CONSUMO"].to_numpy(float))

    return {
        "ars": ars,
        "conv": conv,
        "venc": venc,
        "precios": precios,
        "fecha_txt": fecha_txt,
        "licit": lic,
        "log_cruces": log_cruces,
    }


def exportar_excel_cruzado(data: dict) -> bytes:
    import openpyxl
    base_file = "Reporte Farmacos (NUEVA).xlsx"
    if os.path.exists(base_file):
        wb = openpyxl.load_workbook(base_file)
    else:
        wb = openpyxl.Workbook()
        for s in ['REPORTE', 'ARSENAL_FARMACIA', 'CONVENIOS_VIGENTES', 'NIVELES_STOCK', 'STOCK CENTRALIZADO', 'ICP', 'FALTAS 1708', 'VENCIMIENTOS', 'LICITACIONES', 'FALTAS', 'PEGAR_DATOS', 'CONSUMOS_X_BODEGA']:
            if s not in wb.sheetnames:
                wb.create_sheet(s)

    if "ARSENAL_FARMACIA" in wb.sheetnames:
        ws = wb["ARSENAL_FARMACIA"]
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).value = None
        ars_df = data["ars"]
        for c_idx, col_name in enumerate(ars_df.columns):
            ws.cell(row=1, column=c_idx + 1, value=col_name)
        for r_idx, row in ars_df.iterrows():
            for c_idx, val in enumerate(row):
                ws.cell(row=r_idx + 2, column=c_idx + 1, value=val if pd.notna(val) else None)

    if "CONVENIOS_VIGENTES" in wb.sheetnames:
        ws = wb["CONVENIOS_VIGENTES"]
        conv_df = data["conv"]
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).value = None
        for c_idx, col_name in enumerate(conv_df.columns):
            ws.cell(row=1, column=c_idx + 1, value=col_name)
        for r_idx, row in conv_df.iterrows():
            for c_idx, val in enumerate(row):
                ws.cell(row=r_idx + 2, column=c_idx + 1, value=val if pd.notna(val) else None)

    if "LICITACIONES" in wb.sheetnames:
        ws = wb["LICITACIONES"]
        lic_df = data["licit"]
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).value = None
        for c_idx, col_name in enumerate(lic_df.columns):
            ws.cell(row=1, column=c_idx + 1, value=col_name)
        for r_idx, row in lic_df.iterrows():
            for c_idx, val in enumerate(row):
                ws.cell(row=r_idx + 2, column=c_idx + 1, value=val if pd.notna(val) else None)

    if "VENCIMIENTOS" in wb.sheetnames:
        ws = wb["VENCIMIENTOS"]
        venc_df = data["venc"]
        for r in range(4, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).value = None
        for r_idx, row in venc_df.iterrows():
            for c_idx, val in enumerate(row):
                ws.cell(row=r_idx + 4, column=c_idx + 1, value=val if pd.notna(val) else None)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def f_pm3(y, h):
    return np.full(h, y[-3:].mean())


def f_mean(y, h):
    return np.full(h, y.mean())


def f_ses(y, h, a=0.5):
    lvl = y[0]
    for v in y[1:]:
        lvl = a * v + (1 - a) * lvl
    return np.full(h, lvl)


def f_holt(y, h, a=0.5, b=0.2, phi=0.9):
    k = min(2, len(y) - 1)
    lvl, tr = y[0], (y[k] - y[0]) / max(k, 1)
    for v in y[1:]:
        prev = lvl
        lvl = a * v + (1 - a) * (prev + phi * tr)
        tr = b * (lvl - prev) + (1 - b) * phi * tr
    return lvl + np.cumsum(phi ** np.arange(1, h + 1)) * tr


def f_lin(y, h):
    x = np.arange(len(y))
    slope, inter = np.polyfit(x, y, 1)
    f = inter + slope * np.arange(len(y), len(y) + h)
    return np.clip(f, 0, 1.5 * y.max())


FUNCS = {
    "Promedio 3 meses": f_pm3,
    "Promedio histórico": f_mean,
    "Suavizamiento exponencial": f_ses,
    "Holt amortiguado": f_holt,
    "Tendencia lineal": f_lin,
}
MODEL_NAMES = list(FUNCS)


@st.cache_data(show_spinner="Calculando proyecciones por fármaco…")
def run_forecast(Y: np.ndarray, months: tuple, h: int):
    days = np.array([calendar.monthrange(y, m)[1] for y, m in months], dtype=float)
    R = Y / days
    n, T = R.shape
    M = len(FUNCS)
    F = np.zeros((n, h))
    SIG = np.zeros((n, h))
    best = np.empty(n, dtype=object)
    wape = np.full(n, np.nan)
    wape_m = np.full((n, M), np.nan)
    fns = list(FUNCS.values())
    for i in range(n):
        y = R[i]
        if y.sum() <= 0:
            best[i] = "Sin consumo"
            continue
        abs_e = np.zeros(M)
        sq_e = np.zeros(M)
        k = 0
        for t in range(MIN_TRAIN, T):
            tr = y[:t]
            for j, fn in enumerate(fns):
                e = max(fn(tr, 1)[0], 0.0) - y[t]
                abs_e[j] += abs(e)
                sq_e[j] += e * e
            k += 1
        j = int(np.argmin(np.round(abs_e, 9)))
        test_sum = y[MIN_TRAIN:].sum()
        if test_sum > 0:
            wape_m[i] = abs_e / test_sum
            wape[i] = wape_m[i, j]
        rmse = np.sqrt(sq_e[j] / max(k, 1))
        F[i] = np.clip(fns[j](y, h), 0, None)
        SIG[i] = rmse * np.sqrt(np.arange(1, h + 1))
        best[i] = MODEL_NAMES[j]
    return {"F": F, "SIG": SIG, "best": best, "wape": wape, "wape_m": wape_m, "days": days}


def demanda_diaria(F, fc_months, start, n_days):
    base = fc_months[0][0] * 12 + fc_months[0][1]
    idx = np.empty(n_days, dtype=int)
    for i in range(n_days):
        d = start + timedelta(days=i)
        idx[i] = min(max(d.year * 12 + d.month - base, 0), F.shape[1] - 1)
    return F[:, idx]


def dias_hasta_quiebre(C, D, avail):
    n = len(avail)
    out = np.full(n, np.inf)
    hit = C >= avail[:, None]
    anyhit = hit.any(axis=1)
    k = hit.argmax(axis=1)
    rows = np.arange(n)
    prev = C[rows, k] - D[rows, k]
    frac = np.divide(avail - prev, D[rows, k], out=np.zeros(n), where=D[rows, k] > 0)
    out[anyhit] = (k + frac)[anyhit]
    out[avail <= 0] = 0.0
    out[D.sum(axis=1) <= 0] = np.inf
    return out


N_DIAS_PROY = 730


@st.cache_data(show_spinner="Leyendo stock de insumos…")
def load_stock_ins(file_bytes: bytes):
    texto, d = leer_con_encabezado(file_bytes)
    d = d.rename(columns={"Código": "codigo", "Producto": "producto", "Unidad Medida": "unidad",
                          "Cantidad": "stock", "Precio": "precio", "Total": "valor_stock"})
    for c in ["stock", "precio", "valor_stock"]:
        d[c] = pd.to_numeric(d[c], errors="coerce") if c in d.columns else np.nan
    if "unidad" not in d.columns:
        d["unidad"] = ""
    d["codigo"] = d["codigo"].astype(str).str.strip()
    d["producto"] = d["producto"].astype(str).str.strip()
    d["unidad"] = d["unidad"].fillna("").astype(str)
    d["stock"] = d["stock"].fillna(0)
    d["valor_stock"] = d["valor_stock"].fillna(d["stock"] * d["precio"].fillna(0))
    d = d.groupby("codigo", as_index=False).agg(
        producto=("producto", "first"), unidad=("unidad", "first"), stock=("stock", "sum"),
        precio=("precio", "mean"), valor_stock=("valor_stock", "sum"))
    m = re.search(r"para el d[ií]a\s*(\d{1,2})/(\d{1,2})/(\d{4})", texto)
    fecha = _fecha_dmy(m.group(1), m.group(2), m.group(3)) if m else None
    bodega = re.sub(r"^\s*Stock\s+en\s+", "", texto.split("Obtenido")[0]).strip()
    return d, {"fecha": fecha, "bodega": bodega}


@st.cache_data(show_spinner="Leyendo salidas de insumos…")
def load_salidas(files: tuple):
    partes, bloques, bodega = [], [], ""
    for k, fb in enumerate(files):
        texto, d = leer_con_encabezado(fb)
        d = d.rename(columns={"Código": "codigo", "Producto": "producto", "Unidad Medida": "unidad",
                              "cantidad": "cantidad", "Cantidad": "cantidad", "Total": "total",
                              "Tipo Movimiento": "tipo_mov", "Tipo Documento": "tipo_doc",
                              "Bodega Origen": "origen", "Bodega Destino": "destino"})
        for c in ["unidad", "tipo_mov", "tipo_doc", "origen", "destino"]:
            if c not in d.columns:
                d[c] = ""
            d[c] = d[c].fillna("").astype(str)
        d["codigo"] = d["codigo"].astype(str).str.strip()
        d["producto"] = d["producto"].astype(str).str.strip()
        d["cantidad"] = pd.to_numeric(d["cantidad"], errors="coerce").fillna(0)
        d["total"] = pd.to_numeric(d["total"], errors="coerce").fillna(0)
        m = re.search(r"desde\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*hasta\s*(\d{1,2})/(\d{1,2})/(\d{4})", texto)
        desde = _fecha_dmy(*m.group(1, 2, 3)) if m else None
        hasta = _fecha_dmy(*m.group(4, 5, 6)) if m else None
        fcol = next((c for c in d.columns if "fecha" in str(c).lower()), None)
        if fcol is not None:
            f = pd.to_datetime(d[fcol], errors="coerce", dayfirst=True)
            d["ym"] = f.dt.year * 12 + f.dt.month - 1
            if (desde is None or hasta is None) and f.notna().any():
                desde, hasta = f.min().date(), f.max().date()
        elif desde and hasta and (desde.year, desde.month) == (hasta.year, hasta.month):
            d["ym"] = float(desde.year * 12 + desde.month - 1)
        else:
            d["ym"] = np.nan
        d["bloque"] = k
        partes.append(d[["codigo", "producto", "unidad", "cantidad", "total", "tipo_mov", "tipo_doc",
                         "origen", "destino", "ym", "bloque"]])
        bloques.append((desde, hasta))
        if not bodega:
            bodega = re.sub(r"^\s*Movimientos\s+en\s+", "", texto.split("Obtenido")[0]).strip()
    mov = pd.concat(partes, ignore_index=True)
    dias = sum((h - d_).days + 1 for d_, h in bloques if d_ and h)
    return mov, {"bloques": bloques, "dias": int(dias), "bodega": bodega}


@st.cache_data(show_spinner=False)
def flags_precio_atipico(df: pd.DataFrame):
    k = (df["producto"].str.upper().str.replace(r"[^A-Z0-9 ]", " ", regex=True)
         .str.split().str[:2].str.join(" "))
    out = pd.Series(False, index=df.index)
    for _, idx in df.groupby(k).groups.items():
        if len(idx) < 3:
            continue
        p = df.loc[idx, "precio"]
        for i in idx:
            otros = p.drop(i)
            otros = otros[otros > 0]
            if len(otros) >= 2 and p[i] > 100000 and p[i] > 15 * otros.median():
                out.loc[i] = True
    return out


def construir_serie(cons, bloques, codigos):
    if cons.empty or cons["ym"].isna().any() or not bloques:
        return None
    cov = {}
    for desde, hasta in bloques:
        if desde is None or hasta is None:
            return None
        d = desde
        while d <= hasta:
            key = d.year * 12 + d.month - 1
            cov[key] = cov.get(key, 0) + 1
            d += timedelta(days=1)
    dim = {k: calendar.monthrange(k // 12, k % 12 + 1)[1] for k in cov}
    usables = sorted(k for k in cov if cov[k] >= 0.8 * dim[k])
    if not usables:
        return None
    run = [usables[-1]]
    for k in reversed(usables[:-1]):
        if k == run[-1] - 1:
            run.append(k)
        else:
            break
    run = sorted(run)
    if len(run) < MIN_TRAIN + 1:
        return None
    piv = cons.groupby(["codigo", "ym"])["cantidad"].sum().unstack(fill_value=0)
    piv.columns = [int(c) for c in piv.columns]
    Y = np.zeros((len(codigos), len(run)))
    for j, k in enumerate(run):
        if k in piv.columns:
            s = piv[k].reindex(codigos).fillna(0).to_numpy(float)
            Y[:, j] = s / min(cov[k], dim[k]) * dim[k]
    months = tuple((k // 12, k % 12 + 1) for k in run)
    return months, Y


def calcular_insumos(stock, mov, meta_m, tipos_sel, dias_per, fecha_stock, u_cr, u_ss, horiz_m, dias_seg):
    sal = mov[mov["tipo_mov"].str.upper().str.contains("SALIDA", na=False)]
    cons = sal[sal["tipo_doc"].isin(tipos_sel)]
    g = cons.groupby("codigo").agg(cons_cant=("cantidad", "sum"), cons_valor=("total", "sum"),
                                   n_mov=("cantidad", "size"))
    nombres = sal.drop_duplicates("codigo").set_index("codigo")[["producto", "unidad"]]
    ins = stock.set_index("codigo")
    extra = nombres.loc[~nombres.index.isin(ins.index)]
    if len(extra):
        extra = extra.assign(stock=0.0, precio=np.nan, valor_stock=0.0)
        ins = pd.concat([ins, extra])
    ins.index.name = "codigo"
    ins = ins.join(g)
    for c in ["cons_cant", "cons_valor", "n_mov"]:
        ins[c] = ins[c].fillna(0)
    tot_c = sal.groupby("codigo")["total"].sum()
    tot_q = sal.groupby("codigo")["cantidad"].sum()
    ins["pu_mov"] = (tot_c / tot_q.replace(0, np.nan)).reindex(ins.index)
    ins["precio_arch"] = ins["precio"]
    ins["precio"] = ins["precio"].where(ins["precio"] > 0, ins["pu_mov"])
    ins = ins.reset_index()

    ins["tasa_dia"] = ins["cons_cant"] / max(dias_per, 1)
    ins["aut_dias"] = np.where(ins["tasa_dia"] > 0, ins["stock"] / ins["tasa_dia"].replace(0, np.nan), np.nan)
    ins["estado"] = np.select(
        [(ins["tasa_dia"] <= 0) & (ins["stock"] > 0), ins["tasa_dia"] <= 0, ins["stock"] <= 0,
         ins["aut_dias"] < u_cr, ins["aut_dias"] < u_ss],
        ["SM", "IN", "QS", "CR", "OK"], "SS")
    ins["regularidad"] = np.select(
        [ins["n_mov"] >= 6, ins["n_mov"] >= 3, ins["n_mov"] >= 1],
        ["Regular (6 o más)", "Ocasional (3 a 5)", "Esporádico (1 o 2)"], "Sin consumo")

    ins["precio_atipico"] = flags_precio_atipico(ins[["producto", "precio"]]).to_numpy()
    r = ins["pu_mov"] / ins["precio_arch"]
    ins["precio_cero"] = ins["precio_arch"].fillna(0) <= 0
    ins["precio_difiere"] = (ins["precio_arch"] > 0) & ((r > 3) | (r < 1 / 3))
    ins["precio_sospechoso"] = ins["precio_atipico"] | ins["precio_cero"] | ins["precio_difiere"]

    cods = list(ins["codigo"])
    serie = construir_serie(cons, meta_m["bloques"], cods)
    H = HORIZONTE_FC
    hastas = [h for _, h in meta_m["bloques"] if h]
    corte = max(hastas) if hastas else fecha_stock
    if serie is not None:
        months, Y = serie
        fc = run_forecast(Y, months, H)
        F, SIG = fc["F"], fc["SIG"]
        ins["modelo"] = fc["best"]
        ins["wape"] = fc["wape"]
        ins["confianza"] = np.select(
            [ins["modelo"] == "Sin consumo", ins["wape"].isna(), ins["wape"] < 0.15, ins["wape"] < 0.35],
            ["Sin consumo", "s/d", "Alta", "Media"], "Baja")
        ym_ = months[-1]
        modo = "mensual"
    else:
        months, Y = None, None
        F = np.repeat(ins["tasa_dia"].to_numpy(float)[:, None], H, axis=1)
        SIG = np.zeros_like(F)
        ins["modelo"] = np.where(ins["tasa_dia"] > 0, "Promedio del período", "Sin consumo")
        ins["wape"] = np.nan
        ins["confianza"] = np.where(ins["tasa_dia"] > 0, "Constante", "Sin consumo")
        ym_ = (corte.year, corte.month)
        modo = "constante"
    fcm = []
    for _ in range(H):
        ym_ = next_month(ym_)
        fcm.append(ym_)
    D = demanda_diaria(F, fcm, fecha_stock, N_DIAS_PROY)
    C = np.cumsum(D, axis=1)
    hd = int(round(horiz_m * 30.4375))
    avail = ins["stock"].to_numpy(float)
    dem_h = C[:, hd - 1]
    ins["demanda_h"] = dem_h
    ins["compra"] = np.ceil(np.clip(dem_h + dem_h / hd * dias_seg - avail, 0, None))
    ins["compra_valor"] = ins["compra"] * ins["precio"]
    q = dias_hasta_quiebre(C, D, avail)
    ins["quiebre_dias"] = q
    ins["quiebre_fecha"] = pd.NaT
    fin = np.isfinite(q)
    ins.loc[fin, "quiebre_fecha"] = pd.Timestamp(fecha_stock) + pd.to_timedelta(q[fin], unit="D")
    ins["estado_icon"] = ins["estado"].map(ESTADO_ICON)
    ins["tasa_fc_dia"] = F[:, :3].mean(axis=1)
    return {"ins": ins, "sal": sal, "cons": cons, "F": F, "SIG": SIG, "fcm": fcm, "C": C, "modo": modo,
            "months": months, "Y": Y, "hd": hd, "corte": corte}


def limpiar_servicio(s):
    return re.sub(r"^.*?HOSP\.,\s*", "", str(s)).strip().title()


SIT_COLOR = {"Vencido en bodega activa": "#B3261E", "En riesgo de pérdida": "#E8890C",
             "Se alcanza a consumir": "#2E7D5B", "En cuarentena": "#8A8A8A"}


@st.cache_data(show_spinner="Leyendo vencimientos…")
def load_vencimientos(file_bytes: bytes):
    texto, d = leer_con_encabezado(file_bytes, clave="Establecimiento")
    d = d.rename(columns={"Bodega": "bodega", "Código": "codigo", "Producto": "producto", "Lote": "lote",
                          "Fecha Vencimiento": "vence", "Cantidad": "cantidad",
                          "Precio Lote": "precio_lote", "Total": "total"})
    for c in ["cantidad", "precio_lote", "total"]:
        d[c] = pd.to_numeric(d[c], errors="coerce").fillna(0)
    for c in ["codigo", "producto", "lote", "bodega"]:
        d[c] = d[c].fillna("").astype(str).str.strip()
    v = pd.to_datetime(d["vence"], format="%d/%m/%Y", errors="coerce")
    falta = v.isna()
    if falta.any():
        v[falta] = pd.to_datetime(d.loc[falta, "vence"], errors="coerce")
    d["vence"] = v
    d = d.dropna(subset=["vence"])[["bodega", "codigo", "producto", "lote", "vence", "cantidad",
                                    "precio_lote", "total"]].reset_index(drop=True)
    m = re.search(r"Obtenido\s*:\s*(\d{1,2})/(\d{1,2})/(\d{4})", texto)
    f = re.search(r"Filtro:\s*(.*)$", texto)
    return d, {"fecha": _fecha_dmy(*m.groups()) if m else None, "filtro": f.group(1).strip() if f else ""}


def riesgo_fefo(lots, tasa_map, ref):
    L = lots.sort_values(["codigo", "vence"]).reset_index(drop=True)
    cuar = L["bodega"].str.upper().str.contains("CUARENTENA", na=False).to_numpy()
    cod = L["codigo"].to_numpy()
    cant = L["cantidad"].to_numpy(float)
    dias = (L["vence"] - ref).dt.days.to_numpy(float)
    t = np.clip(dias, 0, None)
    rate = L["codigo"].map(tasa_map).fillna(0).to_numpy(float)
    riesgo = np.zeros(len(L))
    prev, usado = None, 0.0
    for k in range(len(L)):
        if cod[k] != prev:
            prev, usado = cod[k], 0.0
        if cuar[k]:
            riesgo[k] = cant[k]
            continue
        c = min(cant[k], max(0.0, rate[k] * t[k] - usado))
        usado += c
        riesgo[k] = cant[k] - c
    L["dias"] = dias.astype(int)
    L["en_riesgo"] = riesgo
    L["valor_riesgo"] = riesgo * L["precio_lote"]
    L["situacion"] = np.select([cuar, dias < 0, riesgo > 0],
                               ["En cuarentena", "Vencido en bodega activa", "En riesgo de pérdida"],
                               "Se alcanza a consumir")
    L["bodega_corta"] = L["bodega"].str.replace(r"^BODEGA\s+", "", regex=True).str.replace(r"\s+HV$", "", regex=True)
    return L


def _resumen_lotes(lr, base):
    lt = lr[lr["codigo"].isin(set(base["codigo"]))]
    act = lt[lt["situacion"] != "En cuarentena"]
    ag = lt.groupby("codigo").agg(
        lotes=("lote", "size"), unidades=("cantidad", "sum"), en_riesgo=("en_riesgo", "sum"),
        valor_riesgo=("valor_riesgo", "sum"),
        vencidos=("situacion", lambda s: int((s == "Vencido en bodega activa").sum())),
        cuarentena=("situacion", lambda s: int((s == "En cuarentena").sum())))
    ag["primer_venc"] = act.groupby("codigo")["vence"].min()
    ag["dias_primero"] = act.groupby("codigo")["dias"].min()
    ag["en_riesgo_activo"] = act[act["situacion"] == "En riesgo de pérdida"].groupby("codigo")["en_riesgo"].sum()
    t = base.merge(ag, left_on="codigo", right_index=True, how="left")
    for c in ["lotes", "unidades", "en_riesgo", "valor_riesgo", "vencidos", "cuarentena", "en_riesgo_activo"]:
        t[c] = t[c].fillna(0)
    t["activos"] = t["lotes"] - t["cuarentena"]
    t["alerta"] = np.select(
        [t["vencidos"] > 0, t["en_riesgo_activo"] > 0, t["activos"] > 0, t["lotes"] > 0],
        ["🔴 Vencido en bodega activa", "🟠 En riesgo de pérdida", "🟢 Se alcanza a consumir",
         "⚪ Solo lotes en cuarentena"], "🟢 Sin lotes que venzan en el período")
    return t


def render_venc_top(lr, meta_v, prod, clave, que, ref, excl_set=None):
    excl_set = excl_set or set()
    st.caption(
        f"Reporte del {ref:%d/%m/%Y}" + (f" ({meta_v['filtro'].lower()})" if meta_v["filtro"] else "")
        + ". La demanda es la proyectada. Se supone que cada producto consume primero el lote que vence antes "
          "y que las bodegas comparten la demanda. Los lotes en cuarentena no se consumen.")
    c1, c2 = st.columns([1, 2])
    with c1:
        n = st.slider(f"Cuántos {que} de mayor demanda", 3, 30, 10, key=f"{clave}_vn")
    with c2:
        crit = st.radio("Ordenar por", ["Unidades demandadas", "Valor de la demanda ($)"], horizontal=True,
                        key=f"{clave}_vcrit")
    p = prod[prod["demanda_h"] > 0].copy()
    if crit.startswith("Valor"):
        p = p[~p["codigo"].isin(excl_set)]
        p["clave_orden"] = p["demanda_h"] * p["precio"]
    else:
        p["clave_orden"] = p["demanda_h"]
    p = p.sort_values("clave_orden", ascending=False).reset_index(drop=True)
    if p.empty:
        st.info("No hay productos con demanda proyectada en la vista actual.")
        return
    p["puesto"] = np.arange(1, len(p) + 1)
    t_all = _resumen_lotes(lr, p)
    t = t_all.head(n)
    con_venc = t_all[t_all["activos"] > 0]

    n_riesgo = int(((t["vencidos"] > 0) | (t["en_riesgo_activo"] > 0)).sum())
    n_lotes = int((t["activos"] > 0).sum())
    if n_riesgo:
        st.error(f"{n_riesgo} de los {len(t)} {que} de mayor demanda tienen lotes vencidos o que no se alcanzarían "
                 "a consumir antes de vencer.")
    elif n_lotes:
        st.warning(f"{n_lotes} de los {len(t)} {que} de mayor demanda tienen lotes por vencer, pero con la demanda "
                   "proyectada se alcanzan a consumir.")
    else:
        st.success(f"Ninguno de los {len(t)} {que} de mayor demanda tiene lotes por vencer en el período del reporte.")
        if len(con_venc):
            f1 = con_venc.iloc[0]
            st.info(f"El de mayor demanda que sí tiene lotes por vencer es {f1['producto']} (puesto {int(f1['puesto'])}). "
                    "Abajo está la lista de los que sí tienen vencimientos.")
    k = st.columns(4)
    with k[0]:
        kpi("Con lotes por vencer", n0(n_lotes), f"de los {len(t)} de mayor demanda", "#3B6EA8")
    with k[1]:
        kpi("En riesgo de pérdida", n0(int((t["en_riesgo_activo"] > 0).sum())), "No se alcanzan a consumir", "#E8890C")
    with k[2]:
        kpi("Vencidos en bodega activa", n0(int((t["vencidos"] > 0).sum())), "Deberían ir a cuarentena", "#B3261E")
    with k[3]:
        kpi("Valor en riesgo", clp(t["valor_riesgo"].sum()), "Unidades que no se consumirían", "#B3261E")

    cfg = {
        "puesto": st.column_config.NumberColumn("N° en demanda", format="%d"), "alerta": "Alerta",
        "producto": "Producto",
        "demanda_h": st.column_config.NumberColumn("Demanda proyectada", format="%.0f"),
        "unidades": st.column_config.NumberColumn("Unidades por vencer", format="%.0f"),
        "lotes": st.column_config.NumberColumn("Lotes", format="%.0f"),
        "primer_venc": st.column_config.DateColumn("Primer vencimiento", format="DD/MM/YYYY"),
        "dias_primero": st.column_config.NumberColumn("Días al primero", format="%.0f"),
        "en_riesgo": st.column_config.NumberColumn("Unidades en riesgo", format="%.0f"),
        "valor_riesgo": st.column_config.NumberColumn("Valor en riesgo ($)", format="%.0f")}
    cols_t = ["puesto", "alerta", "producto", "demanda_h", "unidades", "lotes", "primer_venc", "dias_primero",
              "en_riesgo", "valor_riesgo"]
    st.markdown(f"**Los {len(t)} {que} de mayor demanda**")
    st.dataframe(t[cols_t], hide_index=True, column_config=cfg)
    st.markdown(f"**Los {que} de mayor demanda que sí tienen lotes por vencer**")
    st.caption("Ordenados por demanda; el número indica su puesto en el ranking completo.")
    st.dataframe(con_venc.head(10)[cols_t], hide_index=True, column_config=cfg)

    mostrar = t if n_lotes else con_venc.head(10)
    lt = lr[lr["codigo"].isin(set(mostrar["codigo"]))].copy()
    if not lt.empty:
        orden = [s_[:42] for s_ in mostrar["producto"]]
        lt = lt.merge(mostrar[["codigo", "producto", "puesto"]].rename(columns={"producto": "prod_top"}),
                      on="codigo", how="left")
        lt["prod_corto"] = lt["prod_top"].str[:42]
        fig = px.scatter(lt, x="vence", y="prod_corto", size="cantidad", color="situacion",
                         color_discrete_map=SIT_COLOR, category_orders={"prod_corto": orden},
                         hover_data={"lote": True, "bodega_corta": True, "cantidad": True, "dias": True,
                                     "prod_corto": False},
                         labels={"vence": "Fecha de vencimiento", "prod_corto": "", "situacion": "Situación"},
                         size_max=30, title="Cuándo vence cada lote (el tamaño es la cantidad)")
        fig.add_shape(type="line", x0=ref.to_pydatetime(), x1=ref.to_pydatetime(), y0=0, y1=1, yref="paper",
                      line=dict(dash="dash", color="#888"))
        fig.add_annotation(x=ref.to_pydatetime(), y=1.02, yref="paper", text="Fecha del reporte", showarrow=False)
        fig.update_layout(height=max(330, 38 * len(orden) + 120), margin=dict(t=60, b=10, l=10, r=10),
                          yaxis=dict(autorange="reversed"), legend=dict(orientation="h", y=-0.15))
        st.plotly_chart(fig)

        with st.expander(f"Lotes de los productos del gráfico ({len(lt)})", expanded=False):
            det = lt.sort_values(["puesto", "vence"])[["puesto", "prod_top", "bodega_corta", "lote", "vence", "dias",
                                                       "cantidad", "en_riesgo", "valor_riesgo", "situacion"]]
            st.dataframe(det, hide_index=True,
                         column_config={
                             "puesto": st.column_config.NumberColumn("N° en demanda", format="%d"),
                             "prod_top": "Producto", "bodega_corta": "Bodega", "lote": "Lote",
                             "vence": st.column_config.DateColumn("Vence", format="DD/MM/YYYY"),
                             "dias": st.column_config.NumberColumn("Días", format="%.0f"),
                             "cantidad": st.column_config.NumberColumn("Cantidad", format="%.0f"),
                             "en_riesgo": st.column_config.NumberColumn("En riesgo", format="%.0f"),
                             "valor_riesgo": st.column_config.NumberColumn("Valor en riesgo ($)", format="%.0f"),
                             "situacion": "Situación"})
            st.download_button("Descargar lotes (CSV)", det.to_csv(index=False).encode("utf-8-sig"),
                               file_name=f"vencimientos_{clave}.csv", mime="text/csv", key=f"dl_venc_{clave}")

    with st.expander(f"Todos los lotes de {que} por vencer ({n0(int(lr['codigo'].isin(set(prod['codigo'])).sum()))})",
                     expanded=False):
        sit_sel = st.multiselect("Situación", list(SIT_COLOR),
                                 default=["Vencido en bodega activa", "En riesgo de pérdida"], key=f"{clave}_vsit")
        nombres = prod.drop_duplicates("codigo").set_index("codigo")["producto"]
        la = lr[lr["codigo"].isin(set(prod["codigo"])) & lr["situacion"].isin(sit_sel)].copy()
        la["producto"] = la["codigo"].map(nombres)
        la = la.sort_values(["dias"])
        st.dataframe(la[["producto", "bodega_corta", "lote", "vence", "dias", "cantidad", "en_riesgo",
                         "valor_riesgo", "situacion"]],
                     hide_index=True,
                     column_config={
                         "producto": "Producto", "bodega_corta": "Bodega", "lote": "Lote",
                         "vence": st.column_config.DateColumn("Vence", format="DD/MM/YYYY"),
                         "dias": st.column_config.NumberColumn("Días", format="%.0f"),
                         "cantidad": st.column_config.NumberColumn("Cantidad", format="%.0f"),
                         "en_riesgo": st.column_config.NumberColumn("En riesgo", format="%.0f"),
                         "valor_riesgo": st.column_config.NumberColumn("Valor en riesgo ($)", format="%.0f"),
                         "situacion": "Situación"})


def render_insumos(stock_file, salidas_files, licit, u_cr, u_ss, horiz_m, dias_seg, venc_file=None):
    if stock_file is None or not salidas_files:
        faltan = []
        if stock_file is None:
            faltan.append("el **stock** de la bodega (reporte «Stock en BODEGA…»)")
        if not salidas_files:
            faltan.append("las **salidas** de la bodega (reporte «Movimientos en BODEGA…»)")
        st.info("Para ver esta pestaña falta subir " + " y ".join(faltan)
                + ". Súbelos en la barra lateral junto con los demás Excel.")
        return
    try:
        stock, meta_s = load_stock_ins(stock_file[1])
        mov, meta_m = load_salidas(tuple(b for _, b in salidas_files))
    except Exception as exc:
        st.error(f"No pude leer los archivos de insumos: {exc}")
        return

    bl = [b for b in meta_m["bloques"] if b[0] and b[1]]
    bl_ord = sorted(bl)
    if any(bl_ord[i][1] >= bl_ord[i + 1][0] for i in range(len(bl_ord) - 1)):
        st.warning("Los archivos de salidas se superponen en fechas. Si subiste el mismo período dos veces, "
                   "el consumo se está contando doble.")

    st.caption(
        f"{meta_s['bodega']}   |   Stock al "
        f"{meta_s['fecha']:%d/%m/%Y}" if meta_s["fecha"] else f"{meta_s['bodega']}"
    )

    with st.expander("Qué cuenta como consumo y otros ajustes", expanded=False):
        tipos_all = sorted(mov["tipo_doc"].unique())
        defecto = [t for t in tipos_all if t in ("SOLICITUD PRODUCTO", "SOLICITUD EXTRAORDINARIA")] or tipos_all
        tipos_sel = st.multiselect(
            "Tipos de salida que cuentan como consumo", tipos_all, default=defecto, key="ins_tipos",
            help="Las solicitudes a servicios son consumo. Los ajustes de inventario salen y llegan a la misma "
                 "bodega, y los traspasos van a otros hospitales; por eso quedan fuera por defecto.")
        dias_per = st.number_input("Días que cubre el archivo de salidas", 1, 1000,
                                   int(meta_m["dias"]) if meta_m["dias"] > 0 else 30, key="ins_dias")
        fecha_stock_i = st.date_input("Fecha del stock de insumos", value=meta_s["fecha"] or date.today(),
                                      key="ins_fecha_stock")
    if not tipos_sel:
        st.warning("Elige al menos un tipo de salida para calcular el consumo.")
        return

    R = calcular_insumos(stock, mov, meta_m, tipos_sel, dias_per, fecha_stock_i, u_cr, u_ss, horiz_m, dias_seg)
    ins, sal, cons, F, SIG, fcm, C, modo = R["ins"], R["sal"], R["cons"], R["F"], R["SIG"], R["fcm"], R["C"], R["modo"]

    tot_cons = ins["cons_valor"].sum()
    sosp = ins[ins["precio_sospechoso"]]
    def_excl = list(sosp[(sosp["cons_valor"] > 0.10 * tot_cons)]["codigo"]) if tot_cons > 0 else []
    etiqueta = {r.codigo: f"{r.producto}" for r in sosp.itertuples()}
    with st.expander(f"Productos con precio sospechoso ({len(sosp)})", expanded=bool(def_excl)):
        st.caption("Un precio muy alto puede inflar todos los montos en pesos. Los productos que elijas aquí "
                   "se dejan fuera de los montos ($), pero siguen contando en los conteos y estados.")
        excl_cods = st.multiselect("Dejar fuera de los montos en pesos", list(sosp["codigo"]),
                                   default=def_excl, format_func=lambda c: etiqueta.get(c, c), key="ins_excl")
    ins["excl"] = ins["codigo"].isin(excl_cods)
    excl_set = set(excl_cods)

    f1, f2 = st.columns([2, 1])
    with f1:
        buscar = st.text_input("Buscar insumo o código", key="ins_buscar")
    with f2:
        reg_opts = ["Regular (6 o más)", "Ocasional (3 a 5)", "Esporádico (1 o 2)", "Sin consumo"]
        reg_sel = st.multiselect("Regularidad del consumo", reg_opts, default=reg_opts, key="ins_reg",
                                 help="Los productos con pocas salidas tienen demanda irregular y su estado "
                                      "puede ser engañoso.")
    v = ins[ins["regularidad"].isin(reg_sel)]
    if buscar.strip():
        t = buscar.strip()
        v = v[v["producto"].str.contains(t, case=False, na=False, regex=False)
              | v["codigo"].str.contains(t, case=False, na=False, regex=False)]
    if v.empty:
        st.warning("Con estos filtros no queda ningún insumo.")
        return
    vm = v[~v["excl"]]

    sub1, sub2, sub3, sub4, sub5, sub6 = st.tabs(
        ["Estado del stock", "Consumo", "Reposición", "Sin movimiento y licitaciones", "Calidad de datos",
         "Vencimientos"])

    with sub1:
        cnt = v["estado"].value_counts().reindex(["QS", "CR", "OK", "SS", "SM", "IN"]).fillna(0).astype(int)
        total_act = int(cnt[["QS", "CR", "OK", "SS", "SM"]].sum())
        desc_estado = {"QS": "sin stock y con consumo", "CR": f"menos de {u_cr} días",
                       "OK": f"{u_cr} a {u_ss - 1} días", "SS": f"{u_ss} días o más"}
        cols = st.columns(4)
        for col, e in zip(cols, ["QS", "CR", "OK", "SS"]):
            with col:
                kpi(ESTADO_TXT[e], n0(cnt[e]), f"{cnt[e] / max(total_act, 1):.1%}  |  {desc_estado[e]}", ESTADO_COLOR[e])
        sm = v[v["estado"] == "SM"]
        cols = st.columns(4)
        with cols[0]:
            kpi("SM (con stock y sin salidas)", n0(cnt["SM"]), f"{clp(sm.loc[~sm['excl'], 'valor_stock'].sum())} detenidos",
                ESTADO_COLOR["SM"])
        with cols[1]:
            kpi("Valor del stock", clp(vm["valor_stock"].sum()), "Stock por precio unitario", "#3B6EA8")
        with cols[2]:
            kpi("Consumo del período", clp(vm["cons_valor"].sum()), f"En {n0(dias_per)} días", "#2E7D5B")
        with cols[3]:
            kpi("IN (sin stock ni salidas)", n0(cnt["IN"]), "Productos sin actividad", ESTADO_COLOR["IN"])
        st.caption(f"Estado calculado con el stock al {fecha_stock_i:%d/%m/%Y} y el consumo diario promedio "
                   f"({n0(dias_per)} días). Las reglas de días son las mismas de fármacos.")

        c1, c2 = st.columns([1, 1.4])
        with c1:
            orden = ["QS", "CR", "OK", "SS", "SM"]
            fig = go.Figure(go.Pie(labels=[ESTADO_TXT[e] for e in orden], values=[cnt[e] for e in orden],
                                   hole=0.55, sort=False, marker=dict(colors=[ESTADO_COLOR[e] for e in orden]),
                                   textinfo="value+percent"))
            fig.update_layout(title="Insumos por estado", height=360, margin=dict(t=50, b=10, l=10, r=10),
                              legend=dict(orientation="h", y=-0.05))
            st.plotly_chart(fig)
        with c2:
            lg = v[v["estado"].isin(["QS", "CR", "OK", "SS", "SM"])].groupby(["regularidad", "estado"]).size().reset_index(name="n")
            fig = px.bar(lg, x="regularidad", y="n", color="estado", color_discrete_map=ESTADO_COLOR,
                         category_orders={"regularidad": reg_opts, "estado": orden},
                         labels={"n": "Insumos", "regularidad": "Regularidad", "estado": "Estado"},
                         title="Estado según regularidad del consumo")
            fig.update_layout(height=360, margin=dict(t=50, b=10, l=10, r=10), legend=dict(orientation="h", y=-0.2))
            st.plotly_chart(fig)

        st.subheader("Los 15 más urgentes")
        st.caption("Insumos en QS o CR con consumo regular u ocasional, ordenados por consumo valorizado.")
        urg = v[v["estado"].isin(["QS", "CR"]) & (v["n_mov"] >= 3)].copy()
        urg["orden_v"] = np.where(urg["excl"], 0, urg["cons_valor"])
        urg = urg.sort_values("orden_v", ascending=False).head(15)
        st.dataframe(
            urg[["estado_icon", "producto", "unidad", "stock", "cons_cant", "n_mov", "tasa_dia", "aut_dias"]],
            hide_index=True,
            column_config={
                "estado_icon": "Estado", "producto": "Insumo", "unidad": "Unidad",
                "stock": st.column_config.NumberColumn("Stock", format="%.0f"),
                "cons_cant": st.column_config.NumberColumn("Consumo del período", format="%.0f"),
                "n_mov": st.column_config.NumberColumn("Salidas", format="%.0f"),
                "tasa_dia": st.column_config.NumberColumn("Consumo diario", format="%.1f"),
                "aut_dias": st.column_config.NumberColumn("Autonomía (días)", format="%.1f")})

        st.subheader("Listado completo")
        est_sel = st.multiselect("Estados a mostrar", ["QS", "CR", "OK", "SS", "SM", "IN"],
                                 default=["QS", "CR", "OK", "SS", "SM"], key="ins_est_lista")
        ords = {e: i for i, e in enumerate(["QS", "CR", "OK", "SS", "SM", "IN"])}
        lst = v[v["estado"].isin(est_sel)].copy()
        lst["o"] = lst["estado"].map(ords)
        lst = lst.sort_values(["o", "aut_dias"], na_position="last")
        vista = lst[["estado_icon", "codigo", "producto", "unidad", "stock", "precio", "valor_stock", "cons_cant",
                     "n_mov", "regularidad", "tasa_dia", "aut_dias"]]
        st.dataframe(
            vista, hide_index=True,
            column_config={
                "estado_icon": "Estado", "codigo": "Código", "producto": "Insumo", "unidad": "Unidad",
                "stock": st.column_config.NumberColumn("Stock", format="%.0f"),
                "precio": st.column_config.NumberColumn("Precio unitario ($)", format="%.0f"),
                "valor_stock": st.column_config.NumberColumn("Valor stock ($)", format="%.0f"),
                "cons_cant": st.column_config.NumberColumn("Consumo del período", format="%.0f"),
                "n_mov": st.column_config.NumberColumn("Salidas", format="%.0f"),
                "regularidad": "Regularidad",
                "tasa_dia": st.column_config.NumberColumn("Consumo diario", format="%.1f"),
                "aut_dias": st.column_config.NumberColumn("Autonomía (días)", format="%.1f")})
        st.download_button("Descargar listado de insumos (CSV)", vista.to_csv(index=False).encode("utf-8-sig"),
                           file_name="estado_insumos.csv", mime="text/csv", key="dl_ins_estado")

    with sub2:
        cods_v = set(v["codigo"])
        cons_v = cons[cons["codigo"].isin(cods_v) & ~cons["codigo"].isin(excl_set)].copy()
        cons_v["servicio"] = cons_v["destino"].map(limpiar_servicio)
        k = st.columns(3)
        with k[0]:
            kpi("Consumo valorizado", clp(cons_v["total"].sum()), f"{n0(cons_v['codigo'].nunique())} productos", "#2E7D5B")
        otros = sal[sal["codigo"].isin(cods_v) & ~sal["codigo"].isin(excl_set)]
        aj = otros[otros["tipo_doc"] == "SOLICITUD AJUSTE"]["total"].sum()
        tr = otros[otros["tipo_doc"] == "TRASPASO"]["total"].sum()
        with k[1]:
            kpi("Ajustes de inventario", clp(aj), "Salen y llegan a la misma bodega", "#E8890C")
        with k[2]:
            kpi("Traspasos a otros hospitales", clp(tr), "Fuera del consumo por defecto", "#3B6EA8")
        if excl_set:
            st.caption(f"Montos sin {len(excl_set)} producto(s) con precio sospechoso (ver el panel de arriba).")

        c1, c2 = st.columns(2)
        with c1:
            sv = cons_v.groupby("servicio")["total"].sum().sort_values(ascending=False).head(12).reset_index()
            fig = px.bar(sv, x="total", y="servicio", orientation="h", color_discrete_sequence=["#2E7D5B"],
                         labels={"total": "Consumo ($)", "servicio": ""}, title="Consumo valorizado por servicio")
            fig.update_layout(height=430, margin=dict(t=50, b=10, l=10, r=10), yaxis=dict(autorange="reversed"))
            st.plotly_chart(fig)
        with c2:
            td = otros.groupby("tipo_doc")["total"].sum().sort_values(ascending=False).reset_index()
            fig = px.bar(td, x="total", y="tipo_doc", orientation="h", color_discrete_sequence=["#3B6EA8"],
                         labels={"total": "Valor de las salidas ($)", "tipo_doc": ""}, title="Salidas por tipo de documento")
            fig.update_layout(height=430, margin=dict(t=50, b=10, l=10, r=10), yaxis=dict(autorange="reversed"))
            st.plotly_chart(fig)

        top = vm.sort_values("cons_valor", ascending=False).head(15)
        fig = px.bar(top, x="cons_valor", y="producto", orientation="h", color="estado",
                     color_discrete_map=ESTADO_COLOR, category_orders={"estado": ["QS", "CR", "OK", "SS", "SM"]},
                     labels={"cons_valor": "Consumo del período ($)", "producto": "", "estado": "Estado"},
                     title="Los 15 insumos con mayor consumo valorizado")
        fig.update_layout(height=480, margin=dict(t=50, b=10, l=10, r=10), yaxis=dict(autorange="reversed"),
                          legend=dict(orientation="h", y=-0.12))
        st.plotly_chart(fig)

        st.subheader("Qué consume cada servicio")
        servs = sorted(cons_v["servicio"].unique())
        if servs:
            sv_sel = st.selectbox("Servicio", servs, key="ins_serv")
            ps = cons_v[cons_v["servicio"] == sv_sel].groupby(["codigo", "producto"]).agg(
                cantidad=("cantidad", "sum"), valor=("total", "sum"), salidas=("cantidad", "size")).reset_index()
            ps = ps.sort_values("valor", ascending=False).head(20)
            st.dataframe(ps[["producto", "cantidad", "valor", "salidas"]], hide_index=True,
                         column_config={"producto": "Insumo",
                                        "cantidad": st.column_config.NumberColumn("Cantidad", format="%.0f"),
                                        "valor": st.column_config.NumberColumn("Valor ($)", format="%.0f"),
                                        "salidas": st.column_config.NumberColumn("Salidas", format="%.0f")})

        st.subheader("Análisis ABC (dónde está el dinero)")
        base_abc = st.radio("Clasificar según", ["Valor del stock", "Consumo valorizado"], horizontal=True,
                            key="ins_abc")
        colv = "valor_stock" if base_abc == "Valor del stock" else "cons_valor"
        ab = vm[vm[colv] > 0].sort_values(colv, ascending=False).reset_index(drop=True)
        if len(ab) >= 3:
            tot = ab[colv].sum()
            share = ab[colv] / tot
            cum = share.cumsum()
            ab["clase"] = np.where(cum - share < 0.80, "A", np.where(cum - share < 0.95, "B", "C"))
            a_, b_, c_ = (ab["clase"] == "A").sum(), (ab["clase"] == "B").sum(), (ab["clase"] == "C").sum()
            k = st.columns(3)
            with k[0]:
                kpi("Clase A", n0(a_), f"Productos que suman el 80% del valor ({a_ / len(ab):.0%} de los productos)", "#B3261E")
            with k[1]:
                kpi("Clase B", n0(b_), "Siguiente 15% del valor", "#E8890C")
            with k[2]:
                kpi("Clase C", n0(c_), "Último 5% del valor", "#8A8A8A")
            fig = go.Figure(go.Scatter(x=(np.arange(len(ab)) + 1) / len(ab) * 100, y=cum * 100, mode="lines",
                                       line=dict(color="#3B6EA8", width=3)))
            fig.add_hline(y=80, line_dash="dash", line_color="#888")
            fig.add_hline(y=95, line_dash="dash", line_color="#888")
            fig.update_layout(height=340, margin=dict(t=20, b=10, l=10, r=10),
                              xaxis_title="% de los productos (de mayor a menor valor)",
                              yaxis_title="% acumulado del valor")
            st.plotly_chart(fig)
            st.dataframe(ab[["clase", "producto", "stock", colv]].head(40), hide_index=True,
                         column_config={"clase": "Clase", "producto": "Insumo",
                                        "stock": st.column_config.NumberColumn("Stock", format="%.0f"),
                                        colv: st.column_config.NumberColumn(base_abc + " ($)", format="%.0f")})
        else:
            st.info("No hay suficientes productos con valor para el análisis ABC.")

    with sub3:
        if modo == "mensual":
            st.caption(f"Proyección con modelos mensuales ({len(R['months'])} meses de historia: "
                       f"{mes_label(R['months'][0])} a {mes_label(R['months'][-1])}).")
        else:
            st.caption("Proyección con el **consumo diario promedio** del período: los movimientos no traen fecha, "
                       "así que no hay tendencia ni rango probable. Con archivos de salidas mes a mes "
                       "(o con una columna de fecha) la pestaña usa los mismos modelos que fármacos.")
        st.caption(f"Compra sugerida = demanda de {horiz_m} meses + {dias_seg} días de seguridad - stock. "
                   "Los meses y días se cambian en la barra lateral (Parámetros de proyección).")
        umb = st.slider("Alertar quiebres dentro de (días)", 7, 180, 30, key="ins_umb")
        solo = st.checkbox("Mostrar solo insumos con compra sugerida", value=True, key="ins_solo")
        solo_reg = st.checkbox("Considerar solo consumo regular y ocasional (3 o más salidas)", value=True, key="ins_solo_reg",
                               help="Los productos con 1 o 2 salidas en el período se piden de vez en cuando; "
                                    "proyectar una compra para ellos suele sobreestimar.")
        vr = v[v["n_mov"] >= 3] if solo_reg else v
        vrm = vr[~vr["excl"]]
        conc = vrm[vrm["compra"] > 0]
        pq = vr[(vr["quiebre_dias"] <= umb) & (vr["estado"].isin(["CR", "OK", "SS"]))]
        k = st.columns(3)
        with k[0]:
            kpi("Insumos con compra sugerida", n0((vr["compra"] > 0).sum()), f"de {n0(len(vr))}", "#3B6EA8")
        with k[1]:
            kpi("Valor estimado de la compra", clp(conc["compra_valor"].sum()), "Solo productos con precio", "#3B6EA8")
        with k[2]:
            kpi(f"Se quedan sin stock en {umb} días", n0(len(pq)), "Sin contar los que ya están en QS", "#E8890C")

        tb = vr.copy()
        if solo:
            tb = tb[tb["compra"] > 0]
        tb["ov"] = np.where(tb["excl"], 0, tb["compra_valor"].fillna(0))
        tb = tb.sort_values("ov", ascending=False)
        vista3 = tb[["estado_icon", "producto", "unidad", "stock", "demanda_h", "compra", "compra_valor",
                     "quiebre_dias", "n_mov", "modelo", "confianza"]]
        st.dataframe(
            vista3, hide_index=True,
            column_config={
                "estado_icon": "Estado", "producto": "Insumo", "unidad": "Unidad",
                "stock": st.column_config.NumberColumn("Stock", format="%.0f"),
                "demanda_h": st.column_config.NumberColumn(f"Demanda {horiz_m} meses", format="%.0f"),
                "compra": st.column_config.NumberColumn("Compra sugerida", format="%.0f"),
                "compra_valor": st.column_config.NumberColumn("Valor compra ($)", format="%.0f"),
                "quiebre_dias": st.column_config.NumberColumn("Días a quiebre", format="%.0f"),
                "n_mov": st.column_config.NumberColumn("Salidas", format="%.0f"),
                "modelo": "Modelo", "confianza": "Confianza"})
        st.download_button("Descargar compras sugeridas de insumos (CSV)",
                           vista3.to_csv(index=False).encode("utf-8-sig"),
                           file_name="compras_insumos.csv", mime="text/csv", key="dl_ins_compra")

        st.subheader("Ficha de un insumo")
        fo = v.assign(o=v["estado"].map(ords)).sort_values(["o", "cons_valor"], ascending=[True, False])
        pos_by_cod = {c: i for i, c in enumerate(ins["codigo"])}
        opciones = list(fo["codigo"])
        cod = st.selectbox("Insumo", opciones, key="ins_ficha",
                           format_func=lambda c: f"{ESTADO_ICON[ins.at[pos_by_cod[c], 'estado']]}  "
                                                 f"{ins.at[pos_by_cod[c], 'producto']}")
        i = pos_by_cod[cod]
        r = ins.loc[i]
        k = st.columns(4)
        with k[0]:
            kpi("Estado", ESTADO_TXT[r["estado"]], f"{n0(r['n_mov'])} salidas en el período", ESTADO_COLOR[r["estado"]])
        with k[1]:
            kpi("Stock", n0(r["stock"]), f"Consumo diario: {r['tasa_dia']:.1f}")
        with k[2]:
            if not np.isfinite(r["quiebre_dias"]):
                kpi("Quiebre proyectado", "Sin quiebre", "Sin consumo proyectado")
            else:
                kpi("Quiebre proyectado", f"{r['quiebre_fecha']:%d/%m/%Y}", f"En {r['quiebre_dias']:.0f} días desde el stock")
        with k[3]:
            kpi("Compra sugerida", n0(r["compra"]), f"Para {horiz_m} meses | {clp(r['compra_valor'])}")

        n_show = int(min(N_DIAS_PROY, max(R["hd"] * 2, 120)))
        fechas = [pd.Timestamp(fecha_stock_i) + pd.Timedelta(days=d) for d in range(n_show + 1)]
        nivel = np.concatenate([[r["stock"]], np.clip(r["stock"] - C[i, :n_show], 0, None)])
        fig = go.Figure(go.Scatter(x=fechas, y=nivel, mode="lines", fill="tozeroy", name="Stock proyectado",
                                   line=dict(color="#3B6EA8", width=3)))
        seg_u = R["C"][i, R["hd"] - 1] / R["hd"] * dias_seg
        fig.add_hline(y=seg_u, line_dash="dash", line_color="#E8890C", annotation_text="Stock de seguridad")
        fig.update_layout(height=340, margin=dict(t=20, b=10, l=10, r=10), yaxis_title="Unidades",
                          title="Stock proyectado (sin ingresos)")
        st.plotly_chart(fig)
        if modo == "mensual":
            hx = [mes_ts(m) for m in R["months"]]
            fx = [mes_ts(m) for m in fcm[:6]]
            dfc = np.array([calendar.monthrange(y, m)[1] for y, m in fcm[:6]], float)
            fy = F[i, :6] * dfc
            flo = np.clip(F[i, :6] - Z80 * SIG[i, :6], 0, None) * dfc
            fhi = (F[i, :6] + Z80 * SIG[i, :6]) * dfc
            fig = go.Figure()
            fig.add_trace(go.Scatter(x=fx + fx[::-1], y=list(fhi) + list(flo[::-1]), fill="toself",
                                     fillcolor="rgba(59,110,168,0.18)", line=dict(width=0),
                                     name="Rango probable (80 %)", hoverinfo="skip"))
            fig.add_trace(go.Scatter(x=hx, y=R["Y"][i], mode="lines+markers", name="Consumo real",
                                     line=dict(color="#2E7D5B", width=3)))
            fig.add_trace(go.Scatter(x=[hx[-1]] + fx, y=[R["Y"][i][-1]] + list(fy), mode="lines+markers",
                                     name="Proyección", line=dict(color="#3B6EA8", width=3, dash="dash")))
            fig.update_layout(height=340, margin=dict(t=20, b=10, l=10, r=10), yaxis_title="Unidades por mes",
                              legend=dict(orientation="h", y=-0.2), hovermode="x unified")
            st.plotly_chart(fig)
        mv = sal[sal["codigo"] == cod].copy()
        mv["destino"] = mv["destino"].map(limpiar_servicio)
        st.markdown("**Salidas de este insumo**")
        st.dataframe(mv[["tipo_doc", "destino", "cantidad", "total"]], hide_index=True,
                     column_config={"tipo_doc": "Tipo de documento", "destino": "Destino",
                                    "cantidad": st.column_config.NumberColumn("Cantidad", format="%.0f"),
                                    "total": st.column_config.NumberColumn("Valor ($)", format="%.0f")})

    with sub4:
        st.subheader("Stock sin movimiento")
        sm_df = v[v["estado"] == "SM"].sort_values("valor_stock", ascending=False)
        kpi("Stock detenido", clp(sm_df.loc[~sm_df["excl"], "valor_stock"].sum()),
            f"{n0(len(sm_df))} productos con stock y sin salidas en el período", ESTADO_COLOR["SM"])
        st.dataframe(sm_df[["producto", "unidad", "stock", "precio", "valor_stock"]], hide_index=True,
                     column_config={"producto": "Insumo", "unidad": "Unidad",
                                    "stock": st.column_config.NumberColumn("Stock", format="%.0f"),
                                    "precio": st.column_config.NumberColumn("Precio unitario ($)", format="%.0f"),
                                    "valor_stock": st.column_config.NumberColumn("Valor ($)", format="%.0f")})

        st.subheader("Licitaciones")
        if licit is None or licit.empty:
            st.info("Para cruzar con licitaciones, sube también `reportes.zip` o el reporte de licitaciones.")
        else:
            lg_ = licit.groupby("Código").agg(
                licitacion=("Licitación", "first"), saldo=("Saldo", "sum"), vcto=("Vencimiento", "min"),
                proveedor=("Proveedor", "first")).reset_index().rename(columns={"Código": "codigo"})
            cz = v.merge(lg_, on="codigo", how="left")
            ref = pd.Timestamp(R["corte"])
            cz["vcto_dias"] = (cz["vcto"] - ref).dt.days
            cz["cob_meses"] = np.where(cz["tasa_dia"] > 0, cz["saldo"] / (cz["tasa_dia"] * 30.4375), np.inf)
            con = cz[cz["licitacion"].notna()]
            urg_ = cz[cz["estado"].isin(["QS", "CR"]) & (cz["n_mov"] >= 3)]
            k = st.columns(3)
            with k[0]:
                kpi("Insumos con licitación", n0(len(con)), f"de {n0(len(cz))} en la vista", "#3B6EA8")
            with k[1]:
                kpi("QS o CR con licitación y saldo", n0((urg_["saldo"] > 0).sum()), "Se pueden reponer con el contrato", "#2E7D5B")
            with k[2]:
                kpi("QS o CR sin licitación o sin saldo", n0(((urg_["saldo"].fillna(0)) <= 0).sum()),
                    "Requieren otra vía de compra", "#B3261E")
            u_v = st.slider("Alertar licitaciones que vencen en menos de (días)", 30, 365, 180, key="ins_u_lic")
            u_c = st.slider("Alertar saldo que cubre menos de (meses)", 1, 12, 3, key="ins_u_cob")
            al = con[(con["vcto_dias"] < u_v) | (con["cob_meses"] < u_c)].sort_values("vcto_dias")
            st.markdown(f"**Licitaciones en alerta: {n0(len(al))}**")
            st.dataframe(al[["estado_icon", "producto", "licitacion", "proveedor", "saldo", "cob_meses", "vcto", "vcto_dias"]],
                         hide_index=True,
                         column_config={
                             "estado_icon": "Estado", "producto": "Insumo", "licitacion": "Licitación",
                             "proveedor": "Proveedor",
                             "saldo": st.column_config.NumberColumn("Saldo", format="%.0f"),
                             "cob_meses": st.column_config.NumberColumn("Cobertura del saldo (meses)", format="%.1f"),
                             "vcto": st.column_config.DateColumn("Vence", format="DD/MM/YYYY"),
                             "vcto_dias": st.column_config.NumberColumn("Días al vencimiento", format="%.0f")})
            st.markdown("**QS o CR sin licitación o sin saldo**")
            sl = urg_[urg_["saldo"].fillna(0) <= 0].sort_values("cons_valor", ascending=False)
            st.dataframe(sl[["estado_icon", "producto", "stock", "cons_cant", "n_mov", "licitacion"]], hide_index=True,
                         column_config={"estado_icon": "Estado", "producto": "Insumo",
                                        "stock": st.column_config.NumberColumn("Stock", format="%.0f"),
                                        "cons_cant": st.column_config.NumberColumn("Consumo del período", format="%.0f"),
                                        "n_mov": st.column_config.NumberColumn("Salidas", format="%.0f"),
                                        "licitacion": "Licitación"})

    with sub5:
        st.caption("Cosas de los archivos que conviene revisar antes de usar los números.")
        sin_stock = sal[~sal["codigo"].isin(set(stock["codigo"]))]["codigo"].nunique()
        k = st.columns(4)
        with k[0]:
            kpi("Precio en cero o negativo", n0(ins["precio_cero"].sum()), "En el archivo de stock", "#E8890C")
        with k[1]:
            kpi("Precio distinto entre archivos", n0(ins["precio_difiere"].sum()), "Salidas vs stock (3 veces o más)", "#E8890C")
        with k[2]:
            kpi("Precio atípico", n0(ins["precio_atipico"].sum()), "Muy sobre productos parecidos", "#B3261E")
        with k[3]:
            kpi("Con salidas pero sin stock registrado", n0(sin_stock), "No están en el archivo de stock", "#3B6EA8")
        rev = ins[ins["precio_sospechoso"]].copy()
        rev["motivo"] = np.select([rev["precio_atipico"], rev["precio_cero"], rev["precio_difiere"]],
                                  ["Precio muy sobre productos parecidos", "Precio en cero o negativo",
                                   "Precio distinto entre stock y salidas"], "")
        rev = rev.sort_values("cons_valor", ascending=False)
        st.dataframe(rev[["producto", "motivo", "precio_arch", "pu_mov", "stock", "cons_valor"]], hide_index=True,
                     column_config={"producto": "Insumo", "motivo": "Motivo",
                                    "precio_arch": st.column_config.NumberColumn("Precio en stock ($)", format="%.0f"),
                                    "pu_mov": st.column_config.NumberColumn("Precio in salidas ($)", format="%.0f"),
                                    "stock": st.column_config.NumberColumn("Stock", format="%.0f"),
                                    "cons_valor": st.column_config.NumberColumn("Consumo ($)", format="%.0f")})
        st.markdown(
            f"- El archivo de salidas cubre **{n0(meta_m['dias'])} días**"
            + (f" ({min(b[0] for b in bl):%d/%m/%Y} a {max(b[1] for b in bl):%d/%m/%Y})." if bl else ".")
            + "\n- El stock está fechado al "
            + (f"**{meta_s['fecha']:%d/%m/%Y}**" if meta_s["fecha"] else "(fecha no informada)")
            + " y no hay ingresos en los archivos, así que no se puede reconstruir el stock de hoy.\n"
            + ("- Los movimientos no traen fecha: la proyección usa el consumo diario promedio."
               if modo == "constante" else "- Hay fechas en los movimientos: la proyección usa modelos mensuales."))

    with sub6:
        if venc_file is None:
            st.info("Sube el reporte de stock por lote («Reporte de Stock… productos que vencen en los próximos N días») "
                    "para ver qué insumos de mayor demanda están por vencer.")
        else:
            try:
                vl, meta_v = load_vencimientos(venc_file[1])
            except Exception as exc:
                st.error(f"No pude leer el vencimientos: {exc}")
            else:
                ref_v = pd.Timestamp(meta_v["fecha"] or date.today())
                lots_i = riesgo_fefo(vl[vl["codigo"].isin(set(ins["codigo"]))],
                                     ins.set_index("codigo")["tasa_fc_dia"], ref_v)
                render_venc_top(lots_i, meta_v, v, "ins", "insumos", ref_v, excl_set)


with st.sidebar:
    st.header("Datos")
    ups = st.file_uploader(
        "Archivos (.zip, .xlsx, .csv)", type=["zip", "xlsx", "csv"], accept_multiple_files=True,
        help="Sube directamente 'reportes.zip' (o los archivos sueltos), junto con el stock y las salidas de IM2.")

archivos = reunir_archivos(ups)
tiene_crudos = hay_archivos_crudos_farmacos(archivos)

with st.sidebar:
    st.caption("Archivos detectados")
    if archivos["farmacos"]:
        st.caption("Fármacos: " + archivos["farmacos"][0] + (" + cruce crudos" if tiene_crudos else ""))
    elif tiene_crudos:
        st.caption("Fármacos: generado automáticamente desde reportes.zip / crudos")
    else:
        st.caption("Fármacos: falta (sube reportes.zip o Reporte_Farmacos)")
    st.caption("Stock de insumos (IM2): " + (archivos["stock_insumos"][0] if archivos["stock_insumos"] else "falta"))
    st.caption("Salidas de insumos (IM2): " + (", ".join(n for n, _ in archivos["salidas_insumos"])
                                               if archivos["salidas_insumos"] else "falta"))
    st.caption("Vencimientos: " + (archivos["vencimientos"][0] if archivos["vencimientos"] else "falta (opcional)"))

if archivos["farmacos"] is None and not tiene_crudos:
    st.title("Fármacos e insumos: estado de stock y demanda")
    st.info("Sube **`reportes.zip`** (o los reportes crudos de fármacos) en la barra lateral para armar y cruzar automáticamente la base de fármacos.")
    st.header("Insumos")
    render_insumos(archivos["stock_insumos"], archivos["salidas_insumos"], None, 30, 360, 3, 20,
                   archivos["vencimientos"])
    st.stop()

if archivos["farmacos"] is not None:
    origen, file_bytes = archivos["farmacos"]
    data_base = load_data(file_bytes)
else:
    origen = "Cruce automático (reportes.zip)"
    data_base = cargar_base_semilla()

data = cruzar_datos_crudos(data_base, archivos)
ars, conv, venc, precios = data["ars"], data["conv"], data["venc"], data["precios"]

if data.get("log_cruces"):
    with st.sidebar:
        st.success("✅ Cruce automático realizado:\n- " + "\n- ".join(data["log_cruces"]))
        st.download_button(
            "📥 Descargar Reporte_Farmacos(NUEVA).xlsx",
            data=exportar_excel_cruzado(data),
            file_name="Reporte Farmacos (NUEVA).xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

cons_cols = sorted(
    [c for c in ars.columns if re.match(r"^[A-Z]{3}-\d{2} CONS$", c) and parse_mes(c)],
    key=lambda c: parse_mes(c),
)
if len(cons_cols) < MIN_TRAIN + 1:
    st.error("No encontré al menos 4 columnas de consumo mensual (formato 'ENE-26 CONS').")
    st.stop()
months = tuple(parse_mes(c) for c in cons_cols)
fc_months = []
ym = months[-1]
for _ in range(HORIZONTE_FC):
    ym = next_month(ym)
    fc_months.append(ym)

try:
    if data["fecha_txt"]:
        dd, mm = data["fecha_txt"]
        y0 = months[-1][0] + (1 if mm < months[-1][1] else 0)
        fecha_def = date(y0, mm, dd)
    else:
        fecha_def = date.today()
except ValueError:
    fecha_def = date.today()

with st.sidebar:
    st.header("Filtros de fármacos")
    excluir_sr = st.toggle(
        "Excluir fármacos sin rotación (SR)",
        value=False,
        help="Activado: misma vista que la tabla dinámica de la hoja REPORTE. "
             "Desactivado: todos los fármacos del arsenal.",
    )
    rot_opts = [r for r in ROT_ORDEN if r in set(ars["ROTACION"].dropna().astype(str))]
    rot_sel = st.multiselect("Rotación", rot_opts, default=rot_opts)
    cen_opts = sorted(set(ars["ESTADO CENABAST"].fillna("Sin Cenabast").astype(str)))
    cen_sel = st.multiselect("Excel Cenabast", cen_opts, default=cen_opts)
    texto = st.text_input("Buscar fármaco o código")

    with st.expander("Parámetros de proyección y compra", expanded=False):
        fecha_stock = st.date_input("Fecha del stock", value=fecha_def,
                                    help="Día al que corresponde el stock del Excel.")
        base_stock = st.radio("Stock a considerar",
                              ["Sin stocklight (como el REPORTE)", "Stock total"], index=0)
        desc_cen = st.checkbox("Sumar entregas Cenabast por llegar", value=True,
                               help="Usa 'CANT ENTREGA CENABAST' como stock que aún no ingresa.")
        horiz_m = st.slider("Meses a cubrir con la compra", 1, 12, 3)
        dias_seg = st.slider("Días de stock de seguridad", 0, 60, 20)

    with st.expander("Reglas de estado (avanzado)", expanded=False):
        u_cr = st.number_input("Crítico si autonomía menor a (días)", 1, 365, 30)
        u_ss = st.number_input("Sobre stock si autonomía desde (días)", 60, 3650, 360)
        rot_recalc = st.checkbox(
            "Recalcular rotación con el mayor consumo de los últimos 3 meses",
            value=False,
            help="En el Excel, MAYOR CONSUMO = MAX(JUN, AGO) y no considera JUL. "
                 "Marca esto para usar los tres meses.",
        )


def num(name, default=0.0):
    if name in ars.columns:
        return pd.to_numeric(ars[name], errors="coerce").fillna(default)
    return pd.Series(default, index=ars.index, dtype=float)


Y = ars[cons_cols].apply(pd.to_numeric, errors="coerce").fillna(0).to_numpy(float)

b = pd.DataFrame({
    "codigo": ars["dato"],
    "descripcion": ars["DESCRIPCION"].astype(str),
    "stock_total": num("STOCK TOTAL"),
    "stock_sin": num("STOCK SIN STOCKLIGHT"),
    "stock_ideal": num("STOCK IDEAL"),
    "consumo_prom": num("CONSUMO PROMEDIO"),
    "consumo_dia": num("CONSUMO DIARIO"),
    "cen_prog": num("CANT PROGRAMADA CENABAST"),
    "cen_falt": num("CANT FALTANTE CENABAST"),
    "cen_ent": num("CANT ENTREGA CENABAST"),
    "estado_cenabast": ars["ESTADO CENABAST"].fillna("Sin Cenabast").astype(str)
    if "ESTADO CENABAST" in ars.columns else "Sin Cenabast",
    "seguimiento": ars["SEGUIMIENTO"] if "SEGUIMIENTO" in ars.columns else "",
})

if "AUTONOMIA DIAS (HOSP - STOCKLIGHT)" in ars.columns:
    b["aut_dias"] = num("AUTONOMIA DIAS (HOSP - STOCKLIGHT)")
else:
    b["aut_dias"] = np.where(b["consumo_dia"] > 0, b["stock_sin"] / b["consumo_dia"].replace(0, np.nan), 0)
b["aut_dias"] = pd.Series(b["aut_dias"]).fillna(0)

b["estado"] = np.select([b["aut_dias"] <= 0, b["aut_dias"] < u_cr, b["aut_dias"] < u_ss],
                        ["QS", "CR", "OK"], "SS")

if rot_recalc or "ROTACION" not in ars.columns:
    b["rotacion"] = rot_from_consumo(Y[:, -3:].max(axis=1))
else:
    b["rotacion"] = ars["ROTACION"].astype(str).values

b["precio"] = b["codigo"].map(precios)

cv = conv.set_index("dato")
b["convenio"] = b["codigo"].map(cv["ID LICITACION"]) if len(cv) else np.nan
b["conv_saldo"] = b["codigo"].map(cv["SALDO"]) if len(cv) else np.nan
b["conv_vcto"] = pd.to_datetime(b["codigo"].map(cv["VCTO"])) if len(cv) else pd.NaT
fecha_stock_ts = pd.Timestamp(fecha_stock)
b["conv_vcto_dias"] = (b["conv_vcto"] - fecha_stock_ts).dt.days

fc = run_forecast(Y, months, HORIZONTE_FC)
F, SIG = fc["F"], fc["SIG"]
days_fc = np.array([calendar.monthrange(y, m)[1] for y, m in fc_months], dtype=float)
b["modelo"] = fc["best"]
b["wape"] = fc["wape"]
b["confianza"] = np.select(
    [b["modelo"] == "Sin consumo", b["wape"].isna(), b["wape"] < 0.15, b["wape"] < 0.35],
    ["Sin consumo", "s/d", "Alta", "Media"], "Baja")

stock_vec = (b["stock_sin"] if base_stock.startswith("Sin") else b["stock_total"]).to_numpy(float)
inc_vec = b["cen_ent"].to_numpy(float) if desc_cen else np.zeros(len(b))
avail = stock_vec + inc_vec
N_DIAS = 730
D = demanda_diaria(F, fc_months, fecha_stock, N_DIAS)
C = np.cumsum(D, axis=1)
hd = int(round(horiz_m * 30.4375))
dem_h = C[:, hd - 1]
seg = dem_h / hd * dias_seg
b["stock_base"] = stock_vec
b["cen_por_llegar"] = inc_vec
b["demanda_h"] = dem_h
b["compra"] = np.ceil(np.clip(dem_h + seg - avail, 0, None))
b["compra_valor"] = b["compra"] * b["precio"]
q = dias_hasta_quiebre(C, D, avail)
b["quiebre_dias"] = q
b["quiebre_fecha"] = pd.NaT
fin = np.isfinite(q)
b.loc[fin, "quiebre_fecha"] = fecha_stock_ts + pd.to_timedelta(q[fin], unit="D")
b["tasa_fc_dia"] = F[:, :3].mean(axis=1)

m1 = fc_months[0]
b["proy_mes1"] = F[:, 0] * days_fc[0]
prog_col = f"PROG {mes_label(m1)}"
b["prog_mes1"] = num(prog_col, np.nan) if prog_col in ars.columns else np.nan


def senal_prog(p, q_):
    if np.isnan(p):
        return "Sin dato"
    if q_ <= 0 and p <= 0:
        return "Sin demanda"
    if p <= 0:
        return "Sin programación"
    if q_ <= 0:
        return "Sobreprogramado"
    r = p / q_
    return "Subprogramado" if r < 0.85 else ("Sobreprogramado" if r > 1.3 else "Alineado")


b["senal_prog"] = [senal_prog(p, q_) for p, q_ in zip(b["prog_mes1"], b["proy_mes1"])]

mask = pd.Series(True, index=b.index)
if excluir_sr:
    mask &= b["rotacion"] != "SR"
if rot_sel:
    mask &= b["rotacion"].isin(rot_sel)
if cen_sel:
    mask &= b["estado_cenabast"].isin(cen_sel)
if texto.strip():
    t = texto.strip()
    mask &= (b["descripcion"].str.contains(t, case=False, na=False, regex=False)
             | b["codigo"].str.contains(t, case=False, na=False, regex=False))
df = b[mask].copy()
df["estado_icon"] = df["estado"].map(ESTADO_ICON)

st.title("Fármacos e insumos: estado de stock y demanda")
st.caption(
    f"Origen: {origen}   |   Stock al {fecha_stock:%d/%m/%Y}   |   "
    f"Consumo real {mes_label(months[0])} a {mes_label(months[-1])}   |   "
    f"Mostrando {n0(len(df))} de {n0(len(b))} fármacos"
    + ("   |   sin rotación (SR) excluidos" if excluir_sr else "")
)

NOMBRES_TABS = ["Estado actual", "Ficha por fármaco", "Proyección y compras", "Alertas", "Insumos"]
if df.empty:
    tabs_v = st.tabs(NOMBRES_TABS)
    with tabs_v[0]:
        st.warning("Con los filtros actuales no queda ningún fármaco. Ajusta los filtros de la izquierda.")
    with tabs_v[4]:
        render_insumos(archivos["stock_insumos"], archivos["salidas_insumos"], data["licit"],
                       u_cr, u_ss, horiz_m, dias_seg, archivos["vencimientos"])
    st.stop()

tab1, tab2, tab3, tab4, tab5 = st.tabs(NOMBRES_TABS)

with tab1:
    cnt = df["estado"].value_counts().reindex(ESTADOS).fillna(0).astype(int)
    total = int(cnt.sum())
    cols = st.columns(4)
    desc_estado = {
        "QS": "sin autonomía",
        "CR": f"menos de {u_cr} días",
        "OK": f"{u_cr} a {u_ss - 1} días",
        "SS": f"{u_ss} días o más",
    }
    for col, e in zip(cols, ESTADOS):
        with col:
            kpi(ESTADO_TXT[e], n0(cnt[e]), f"{cnt[e] / total:.1%} del total  |  {desc_estado[e]}", ESTADO_COLOR[e])

    c1, c2 = st.columns([1, 1.4])
    with c1:
        fig = go.Figure(go.Pie(
            labels=[ESTADO_TXT[e] for e in ESTADOS], values=[cnt[e] for e in ESTADOS],
            hole=0.55, sort=False, marker=dict(colors=[ESTADO_COLOR[e] for e in ESTADOS]),
            textinfo="value+percent"))
        fig.update_layout(title="Fármacos por estado", height=360, margin=dict(t=50, b=10, l=10, r=10),
                          legend=dict(orientation="h", y=-0.05))
        st.plotly_chart(fig)
    with c2:
        ct = pd.crosstab(df["estado"], df["rotacion"]).reindex(ESTADOS).fillna(0).astype(int)
        ct = ct[[r for r in ROT_ORDEN if r in ct.columns]]
        long = ct.reset_index().melt(id_vars="estado", var_name="rotacion", value_name="n")
        fig = px.bar(long, x="rotacion", y="n", color="estado", color_discrete_map=ESTADO_COLOR,
                     category_orders={"estado": ESTADOS, "rotacion": list(ct.columns)},
                     labels={"n": "Fármacos", "rotacion": "Rotación", "estado": "Estado"},
                     title="Estado según rotación")
        fig.update_layout(height=360, margin=dict(t=50, b=10, l=10, r=10), legend=dict(orientation="h", y=-0.2))
        st.plotly_chart(fig)

    with st.expander("Tabla estado por rotación (igual a la tabla dinámica del REPORTE)", expanded=False):
        ct_t = ct.copy()
        ct_t["Total"] = ct_t.sum(axis=1)
        ct_t.loc["Total"] = ct_t.sum()
        st.dataframe(ct_t)

    c3, c4 = st.columns(2)
    with c3:
        h = df[df["estado"] != "QS"].copy()
        h["aut_cap"] = np.minimum(h["aut_dias"], 720)
        fig = px.histogram(h, x="aut_cap", color="estado", nbins=48, color_discrete_map=ESTADO_COLOR,
                           category_orders={"estado": ESTADOS},
                           labels={"aut_cap": "Autonomía (días)", "estado": "Estado"},
                           title="Distribución de la autonomía (sin QS)")
        fig.add_vline(x=u_cr, line_dash="dash", line_color="#888")
        fig.add_vline(x=u_ss, line_dash="dash", line_color="#888")
        fig.update_layout(height=340, margin=dict(t=50, b=10, l=10, r=10), legend=dict(orientation="h", y=-0.25),
                          yaxis_title="Fármacos")
        st.plotly_chart(fig)
        st.caption("Los valores sobre 720 días se agrupan en el último tramo.")
    with c4:
        fig = px.bar(df.groupby(["estado_cenabast", "estado"]).size().reset_index(name="n"),
                     x="estado_cenabast", y="n", color="estado", color_discrete_map=ESTADO_COLOR,
                     category_orders={"estado": ESTADOS},
                     labels={"n": "Fármacos", "estado_cenabast": "Estado Cenabast", "estado": "Estado"},
                     title="Estado del stock según estado Cenabast")
        fig.update_layout(height=340, margin=dict(t=50, b=10, l=10, r=10), legend=dict(orientation="h", y=-0.25))
        st.plotly_chart(fig)

    st.subheader("Los 15 más urgentes")
    st.caption("Fármacos en QS o CR con consumo, ordenados por consumo mensual promedio.")
    urg = df[df["estado"].isin(["QS", "CR"]) & (df["consumo_prom"] > 0)].sort_values(
        "consumo_prom", ascending=False).head(15)
    st.dataframe(
        urg[["estado_icon", "descripcion", "rotacion", "stock_sin", "consumo_dia", "aut_dias",
             "estado_cenabast", "cen_por_llegar", "seguimiento"]],
        hide_index=True,
        column_config={
            "estado_icon": "Estado", "descripcion": "Fármaco", "rotacion": "Rotación",
            "stock_sin": st.column_config.NumberColumn("Stock sin stocklight", format="%.0f"),
            "consumo_dia": st.column_config.NumberColumn("Consumo diario", format="%.0f"),
            "aut_dias": st.column_config.NumberColumn("Autonomía (días)", format="%.1f"),
            "estado_cenabast": "Cenabast",
            "cen_por_llegar": st.column_config.NumberColumn("Entrega Cenabast", format="%.0f"),
            "seguimiento": "Seguimiento",
        })

    st.subheader("Listado completo")
    est_sel = st.multiselect("Estados a mostrar", ESTADOS, default=ESTADOS, key="est_tab1")
    lst = df[df["estado"].isin(est_sel)].copy()
    lst["orden"] = lst["estado"].map({e: i for i, e in enumerate(ESTADOS)})
    lst = lst.sort_values(["orden", "aut_dias"])
    vista = lst[["estado_icon", "codigo", "descripcion", "rotacion", "stock_sin", "stock_total",
                 "consumo_prom", "consumo_dia", "aut_dias", "estado_cenabast", "convenio",
                 "conv_vcto_dias", "seguimiento"]]
    st.dataframe(
        vista, hide_index=True,
        column_config={
            "estado_icon": "Estado", "codigo": "Código", "descripcion": "Fármaco", "rotacion": "Rotación",
            "stock_sin": st.column_config.NumberColumn("Stock sin stocklight", format="%.0f"),
            "stock_total": st.column_config.NumberColumn("Stock total", format="%.0f"),
            "consumo_prom": st.column_config.NumberColumn("Consumo mensual prom.", format="%.0f"),
            "consumo_dia": st.column_config.NumberColumn("Consumo diario", format="%.0f"),
            "aut_dias": st.column_config.NumberColumn("Autonomía (días)", format="%.1f"),
            "estado_cenabast": "Cenabast", "convenio": "Convenio",
            "conv_vcto_dias": st.column_config.NumberColumn("Convenio vence en (días)", format="%.0f"),
            "seguimiento": "Seguimiento",
        })
    st.download_button("Descargar listado (CSV)", vista.to_csv(index=False).encode("utf-8-sig"),
                       file_name="estado_farmacos.csv", mime="text/csv")

with tab2:
    orden_sel = df.assign(o=df["estado"].map({e: i for i, e in enumerate(ESTADOS)})).sort_values(
        ["o", "consumo_prom"], ascending=[True, False])
    opciones = list(orden_sel.index)
    idx = st.selectbox("Fármaco", opciones,
                       format_func=lambda i: f"{ESTADO_ICON[b.at[i, 'estado']]}  {b.at[i, 'descripcion']}",
                       help="Escribe para buscar. Primero aparecen los QS y críticos con más consumo.")
    r = b.loc[idx]
    e = r["estado"]

    k = st.columns(5)
    with k[0]:
        kpi("Estado", ESTADO_TXT[e], f"Rotación {r['rotacion']}", ESTADO_COLOR[e])
    with k[1]:
        kpi("Autonomía actual", f"{r['aut_dias']:.0f} días", f"Stock sin stocklight: {n0(r['stock_sin'])}")
    with k[2]:
        kpi("Consumo diario", n0(r["consumo_dia"]), f"Mensual promedio: {n0(r['consumo_prom'])}")
    with k[3]:
        qd = r["quiebre_dias"]
        if not np.isfinite(qd):
            kpi("Quiebre proyectado", "Sin quiebre", "En los próximos 2 años")
        else:
            kpi("Quiebre proyectado", f"{r['quiebre_fecha']:%d/%m/%Y}", f"En {qd:.0f} días")
    with k[4]:
        kpi("Compra sugerida", n0(r["compra"]), f"Para {horiz_m} meses | {clp(r['compra_valor'])}")

    st.subheader("Consumo mensual: real y proyectado")
    hist_x = [mes_ts(m) for m in months]
    nfc = st.slider("Meses a proyectar en el gráfico", 1, HORIZONTE_FC, 6, key="nfc_ficha")
    fx = [mes_ts(m) for m in fc_months[:nfc]]
    fy = F[idx, :nfc] * days_fc[:nfc]
    flo = np.clip(F[idx, :nfc] - Z80 * SIG[idx, :nfc], 0, None) * days_fc[:nfc]
    fhi = (F[idx, :nfc] + Z80 * SIG[idx, :nfc]) * days_fc[:nfc]
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=fx + fx[::-1], y=list(fhi) + list(flo[::-1]), fill="toself",
                             fillcolor="rgba(59,110,168,0.18)", line=dict(width=0),
                             name="Rango probable (80 %)", hoverinfo="skip"))
    fig.add_trace(go.Scatter(x=hist_x, y=Y[idx], mode="lines+markers", name="Consumo real",
                             line=dict(color="#2E7D5B", width=3)))
    fig.add_trace(go.Scatter(x=[hist_x[-1]] + fx, y=[Y[idx][-1]] + list(fy), mode="lines+markers",
                             name="Proyección", line=dict(color="#3B6EA8", width=3, dash="dash")))
    prog_x, prog_y = [], []
    for ym_ in list(months) + fc_months:
        cname = f"PROG {mes_label(ym_)}"
        if cname in ars.columns:
            v = ars.at[idx, cname]
            if pd.notna(v) and v > 0:
                prog_x.append(mes_ts(ym_))
                prog_y.append(v)
    if prog_x:
        fig.add_trace(go.Scatter(x=prog_x, y=prog_y, mode="markers", name="Programado",
                                 marker=dict(symbol="diamond", size=11, color="#E8890C")))
    fig.update_layout(height=400, margin=dict(t=20, b=10, l=10, r=10), yaxis_title="Unidades por mes",
                      legend=dict(orientation="h", y=-0.15), hovermode="x unified")
    st.plotly_chart(fig)
    st.caption(f"Modelo elegido: {r['modelo']}   |   Confianza: {r['confianza']}"
               + (f"   |   Error histórico del modelo (WAPE): {r['wape']:.0%}" if pd.notna(r["wape"]) else ""))

    c1, c2 = st.columns(2)
    with c1:
        st.subheader("Stock por ubicación")
        ubic = {
            "Farmacia AT abierta": "STOCK AT ABIERTA", "Farmacia AT cerrada": "STOCK AT CERRADA",
            "Stocklight abierta": "STOCKLIGTH ABIERTA", "Stocklight cerrada": "STOCKTIGTH CERADA",
            "Botiquín urgencia": "BOTIQUIN URGENCIA", "Bodega CEM": "BODEGA CEM",
            "Bodega fármacos": "FARMACOS", "Bodega fármacos 2": "FARMACOS 2",
        }
        ud = pd.DataFrame({"Ubicación": [k_ for k_, v in ubic.items() if v in ars.columns],
                           "Unidades": [ars.at[idx, v] for v in ubic.values() if v in ars.columns]})
        ud["Unidades"] = pd.to_numeric(ud["Unidades"], errors="coerce").fillna(0)
        fig = px.bar(ud, x="Unidades", y="Ubicación", orientation="h", color_discrete_sequence=["#3B6EA8"])
        fig.update_layout(height=330, margin=dict(t=10, b=10, l=10, r=10), yaxis=dict(autorange="reversed"))
        st.plotly_chart(fig)
    with c2:
        st.subheader("Qué modelo acierta más")
        wm = fc["wape_m"][idx]
        if np.all(np.isnan(wm)):
            st.info("No hay consumo suficiente para evaluar modelos.")
        else:
            md = pd.DataFrame({"Modelo": MODEL_NAMES, "Error (WAPE)": wm})
            md["Elegido"] = np.where(md["Modelo"] == r["modelo"], "Sí", "")
            fig = px.bar(md, x="Error (WAPE)", y="Modelo", orientation="h", color="Elegido",
                         color_discrete_map={"Sí": "#3B6EA8", "": "#B5B5B5"})
            fig.update_layout(height=330, margin=dict(t=10, b=10, l=10, r=10), showlegend=False,
                              yaxis=dict(autorange="reversed"), xaxis_tickformat=".0%")
            st.plotly_chart(fig)
            st.caption("Cada modelo intentó adivinar los últimos meses con datos anteriores. Menor error es mejor.")

    d1, d2, d3 = st.columns(3)
    with d1:
        st.markdown("**Cenabast**")
        st.write(f"Estado: {r['estado_cenabast']}")
        st.write(f"Programado: {n0(r['cen_prog'])}")
        st.write(f"Entrega aprobada: {n0(r['cen_ent'])}")
        st.write(f"Faltante: {n0(r['cen_falt'])}")
    with d2:
        st.markdown("**Convenio**")
        if pd.isna(r["convenio"]):
            st.write("Sin convenio vigente")
        else:
            st.write(f"Licitación: {r['convenio']}")
            st.write(f"Saldo: {n0(r['conv_saldo'])}")
            vct = "sin fecha" if pd.isna(r["conv_vcto"]) else f"{r['conv_vcto']:%d/%m/%Y} ({n0(r['conv_vcto_dias'])} días)"
            st.write(f"Vence: {vct}")
    with d3:
        st.markdown("**Seguimiento**")
        st.write(r["seguimiento"] if pd.notna(r["seguimiento"]) and str(r["seguimiento"]).strip() else "Sin comentarios")
        st.write(f"Precio unitario estimado: {clp(r['precio'])}")

with tab3:
    st.caption(
        f"Compra sugerida = demanda proyectada de {horiz_m} meses + {dias_seg} días de seguridad "
        f"- stock ({'sin stocklight' if base_stock.startswith('Sin') else 'total'})"
        + (" - entregas Cenabast por llegar" if desc_cen else "")
        + ". Los parámetros se cambian en la barra lateral."
    )
    horiz_q = st.slider("Alertar quiebres dentro de (días)", 7, 180, 30, key="q_dias_t3")
    con_compra = df[df["compra"] > 0]
    quiebres = df[(df["quiebre_dias"] <= horiz_q) & (df["estado"] != "QS")]
    k = st.columns(4)
    with k[0]:
        kpi("Fármacos con compra sugerida", n0(len(con_compra)), f"de {n0(len(df))}", "#3B6EA8")
    with k[1]:
        kpi("Valor estimado de la compra", clp(con_compra["compra_valor"].sum()),
            "Solo fármacos con precio conocido", "#3B6EA8")
    with k[2]:
        kpi(f"Quiebres en {horiz_q} días", n0(len(quiebres)), "Sin contar los que ya están en QS", "#E8890C")
    with k[3]:
        sub = df[df["senal_prog"].isin(["Subprogramado", "Sin programación"])]
        kpi(f"Subprogramados en {mes_label(m1)}", n0(len(sub)), "Programación menor a la demanda proyectada", "#B3261E")

    st.subheader("Demanda total valorizada")
    dv = df[df["precio"].notna()]
    ii = dv.index.to_numpy()
    pr = dv["precio"].to_numpy(float)
    hist_val = (Y[ii] * pr[:, None]).sum(axis=0)
    nf = st.slider("Meses a proyectar", 1, HORIZONTE_FC, 6, key="nfc_total")
    fc_val = (F[ii, :nf] * days_fc[:nf] * pr[:, None]).sum(axis=0)
    sig_val = np.sqrt((((SIG[ii, :nf] * days_fc[:nf]) * pr[:, None]) ** 2).sum(axis=0))
    fx = [mes_ts(m) for m in fc_months[:nf]]
    hx = [mes_ts(m) for m in months]
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=fx + fx[::-1], y=list(fc_val + Z80 * sig_val) + list(np.clip(fc_val - Z80 * sig_val, 0, None)[::-1]),
                             fill="toself", fillcolor="rgba(59,110,168,0.18)", line=dict(width=0),
                             name="Rango probable (80 %)", hoverinfo="skip"))
    fig.add_trace(go.Scatter(x=hx, y=hist_val, mode="lines+markers", name="Consumo real",
                             line=dict(color="#2E7D5B", width=3)))
    fig.add_trace(go.Scatter(x=[hx[-1]] + fx, y=[hist_val[-1]] + list(fc_val), mode="lines+markers",
                             name="Proyección", line=dict(color="#3B6EA8", width=3, dash="dash")))
    fig.update_layout(height=380, margin=dict(t=20, b=10, l=10, r=10), yaxis_title="Pesos por mes",
                      legend=dict(orientation="h", y=-0.15), hovermode="x unified")
    st.plotly_chart(fig)
    st.caption("Consumo multiplicado por el precio unitario estimado desde STOCK CENTRALIZADO. "
               f"Cubre {n0(len(dv))} de {n0(len(df))} fármacos (los demás no tienen precio).")

    top = df.assign(valor_h=df["demanda_h"] * df["precio"]).sort_values("valor_h", ascending=False).head(15)
    fig = px.bar(top, x="valor_h", y="descripcion", orientation="h", color="estado",
                 color_discrete_map=ESTADO_COLOR, category_orders={"estado": ESTADOS},
                 labels={"valor_h": f"Demanda proyectada {horiz_m} meses ($)", "descripcion": "", "estado": "Estado"},
                 title=f"Los 15 fármacos con mayor demanda valorizada en los próximos {horiz_m} meses")
    fig.update_layout(height=470, margin=dict(t=50, b=10, l=10, r=10), yaxis=dict(autorange="reversed"),
                      legend=dict(orientation="h", y=-0.12))
    st.plotly_chart(fig)

    st.subheader("Compra sugerida y programación")
    solo_compra = st.checkbox("Mostrar solo fármacos con compra sugerida", value=True)
    señales = ["Subprogramado", "Sin programación", "Alineado", "Sobreprogramado", "Sin demanda", "Sin dato"]
    sel_sen = st.multiselect(f"Programación {mes_label(m1)}", señales, default=señales)
    tb = df[df["senal_prog"].isin(sel_sen)]
    if solo_compra:
        tb = tb[tb["compra"] > 0]
    tb = tb.sort_values("compra_valor", ascending=False, na_position="last")
    vista3 = tb[["estado_icon", "descripcion", "stock_base", "cen_por_llegar", "demanda_h", "compra",
                 "compra_valor", "proy_mes1", "prog_mes1", "senal_prog", "quiebre_dias", "modelo", "confianza"]]
    st.dataframe(
        vista3, hide_index=True,
        column_config={
            "estado_icon": "Estado", "descripcion": "Fármaco",
            "stock_base": st.column_config.NumberColumn("Stock considerado", format="%.0f"),
            "cen_por_llegar": st.column_config.NumberColumn("Cenabast por llegar", format="%.0f"),
            "demanda_h": st.column_config.NumberColumn(f"Demanda {horiz_m} meses", format="%.0f"),
            "compra": st.column_config.NumberColumn("Compra sugerida", format="%.0f"),
            "compra_valor": st.column_config.NumberColumn("Valor compra ($)", format="%.0f"),
            "proy_mes1": st.column_config.NumberColumn(f"Proyectado {mes_label(m1)}", format="%.0f"),
            "prog_mes1": st.column_config.NumberColumn(f"Programado {mes_label(m1)}", format="%.0f"),
            "senal_prog": "Programación",
            "quiebre_dias": st.column_config.NumberColumn("Días a quiebre", format="%.0f"),
            "modelo": "Modelo", "confianza": "Confianza",
        })
    st.download_button("Descargar compras sugeridas (CSV)", vista3.to_csv(index=False).encode("utf-8-sig"),
                       file_name="compras_sugeridas.csv", mime="text/csv")

    with st.expander("Cómo se calcula la proyección", expanded=False):
        st.markdown(
            "- El consumo mensual se pasa a **consumo diario** (dividiendo por los días de cada mes) para que "
            "febrero no parezca una baja.\n"
            "- Para cada fármaco se prueban 5 modelos simples: promedio de 3 meses, promedio histórico, "
            "suavizamiento exponencial, Holt amortiguado y tendencia lineal.\n"
            "- Cada modelo intenta adivinar mes a mes los últimos meses usando solo los anteriores. "
            "Se elige el de menor error (WAPE) y ese error define el **rango probable del 80 %**.\n"
            "- El **quiebre proyectado** recorre día a día la demanda proyectada hasta agotar el stock disponible."
        )

# ===========================================================================
# TAB 4: Alertas
# ===========================================================================
with tab4:
    # 1. Quiebres próximos
    st.subheader("Quiebres próximos")
    ya_qs = df[(df["estado"] == "QS") & (df["tasa_fc_dia"] > 0)]
    umbral_q = st.slider("Mostrar fármacos que se quedan sin stock en menos de (días)", 7, 180, 30, key="q_dias_t4")
    pq = df[(df["quiebre_dias"] <= umbral_q) & (df["estado"] != "QS")].sort_values("quiebre_dias")
    a1, a2 = st.columns(2)
    with a1:
        kpi("Se quedan sin stock pronto", n0(len(pq)), f"En menos de {umbral_q} días", "#E8890C")
    with a2:
        kpi("Ya en QS pero con demanda proyectada", n0(len(ya_qs)), "Hay que reponer", "#B3261E")
    st.dataframe(
        pq[["estado_icon", "descripcion", "stock_base", "cen_por_llegar", "tasa_fc_dia", "quiebre_dias",
            "quiebre_fecha", "estado_cenabast", "compra"]],
        hide_index=True,
        column_config={
            "estado_icon": "Estado", "descripcion": "Fármaco",
            "stock_base": st.column_config.NumberColumn("Stock considerado", format="%.0f"),
            "cen_por_llegar": st.column_config.NumberColumn("Cenabast por llegar", format="%.0f"),
            "tasa_fc_dia": st.column_config.NumberColumn("Demanda diaria proyectada", format="%.0f"),
            "quiebre_dias": st.column_config.NumberColumn("Días a quiebre", format="%.0f"),
            "quiebre_fecha": st.column_config.DateColumn("Fecha de quiebre", format="DD/MM/YYYY"),
            "estado_cenabast": "Cenabast",
            "compra": st.column_config.NumberColumn("Compra sugerida", format="%.0f"),
        })
    with st.expander("Fármacos ya en QS con demanda proyectada"):
        st.dataframe(
            ya_qs.sort_values("tasa_fc_dia", ascending=False)[
                ["descripcion", "rotacion", "tasa_fc_dia", "estado_cenabast", "cen_por_llegar", "compra", "seguimiento"]],
            hide_index=True,
            column_config={
                "descripcion": "Fármaco", "rotacion": "Rotación",
                "tasa_fc_dia": st.column_config.NumberColumn("Demanda diaria proyectada", format="%.1f"),
                "estado_cenabast": "Cenabast",
                "cen_por_llegar": st.column_config.NumberColumn("Cenabast por llegar", format="%.0f"),
                "compra": st.column_config.NumberColumn("Compra sugerida", format="%.0f"),
                "seguimiento": "Seguimiento"})

    # 2. Cenabast en riesgo
    st.subheader("Cenabast con problemas y stock bajo")
    riesgo = df[df["estado_cenabast"].isin(CEN_RIESGO)]
    cz = riesgo[riesgo["estado"].isin(["QS", "CR"])].sort_values("consumo_prom", ascending=False)
    b1, b2 = st.columns([1, 2])
    with b1:
        kpi("Cenabast FALTANTE, SUSPENDIDO o ELIMINADO", n0(len(riesgo)), f"{n0(len(cz))} están en QS o CR", "#B3261E")
        if len(riesgo):
            mt = pd.crosstab(riesgo["estado_cenabast"], riesgo["estado"]).reindex(columns=ESTADOS).fillna(0).astype(int)
            st.dataframe(mt)
    with b2:
        st.dataframe(
            cz[["estado_icon", "descripcion", "estado_cenabast", "cen_prog", "cen_falt", "stock_sin",
                "aut_dias", "consumo_prom"]],
            hide_index=True,
            column_config={
                "estado_icon": "Estado", "descripcion": "Fármaco", "estado_cenabast": "Cenabast",
                "cen_prog": st.column_config.NumberColumn("Programado", format="%.0f"),
                "cen_falt": st.column_config.NumberColumn("Faltante", format="%.0f"),
                "stock_sin": st.column_config.NumberColumn("Stock sin stocklight", format="%.0f"),
                "aut_dias": st.column_config.NumberColumn("Autonomía (días)", format="%.1f"),
                "consumo_prom": st.column_config.NumberColumn("Consumo mensual prom.", format="%.0f")})

    # 3. Sobre stock
    st.subheader("Sobre stock y dinero inmovilizado")
    ss = df[df["estado"] == "SS"].copy()
    ss["exceso_u"] = np.clip(ss["stock_sin"] - ss["stock_ideal"], 0, None)
    ss["exceso_val"] = ss["exceso_u"] * ss["precio"]
    ss["stock_val"] = ss["stock_sin"] * ss["precio"]
    c1, c2, c3 = st.columns(3)
    with c1:
        kpi("Fármacos en sobre stock", n0(len(ss)), f"Autonomía de {u_ss} días o más", ESTADO_COLOR["SS"])
    with c2:
        kpi("Valor del stock en SS", clp(ss["stock_val"].sum()), "Stock sin stocklight por precio", ESTADO_COLOR["SS"])
    with c3:
        kpi("Exceso sobre el stock ideal", clp(ss["exceso_val"].sum()), "Stock ideal = 60 días de consumo", ESTADO_COLOR["SS"])
    top_ss = ss.sort_values("exceso_val", ascending=False).head(15)
    fig = px.bar(top_ss, x="exceso_val", y="descripcion", orientation="h",
                 color_discrete_sequence=[ESTADO_COLOR["SS"]],
                 labels={"exceso_val": "Exceso valorizado ($)", "descripcion": ""},
                 title="Los 15 con mayor exceso valorizado")
    fig.update_layout(height=450, margin=dict(t=50, b=10, l=10, r=10), yaxis=dict(autorange="reversed"))
    st.plotly_chart(fig)
    st.dataframe(
        ss.sort_values("exceso_val", ascending=False)[
            ["descripcion", "rotacion", "stock_sin", "stock_ideal", "aut_dias", "exceso_u", "exceso_val"]],
        hide_index=True,
        column_config={
            "descripcion": "Fármaco", "rotacion": "Rotación",
            "stock_sin": st.column_config.NumberColumn("Stock sin stocklight", format="%.0f"),
            "stock_ideal": st.column_config.NumberColumn("Stock ideal", format="%.0f"),
            "aut_dias": st.column_config.NumberColumn("Autonomía (días)", format="%.0f"),
            "exceso_u": st.column_config.NumberColumn("Exceso (unidades)", format="%.0f"),
            "exceso_val": st.column_config.NumberColumn("Exceso ($)", format="%.0f")})

    # 4. Convenios
    st.subheader("Convenios por vencer o con poco saldo")
    u_conv = st.slider("Alertar convenios que vencen en menos de (días)", 30, 365, 180, key="u_conv")
    u_cob = st.slider("Alertar saldo que cubre menos de (meses)", 1, 12, 3, key="u_cob")
    cv_df = df[df["convenio"].notna()].copy()
    cv_df["cob_meses"] = np.where(cv_df["tasa_fc_dia"] > 0,
                                  cv_df["conv_saldo"] / (cv_df["tasa_fc_dia"] * 30.4375), np.inf)
    al = cv_df[(cv_df["conv_vcto_dias"] < u_conv) | (cv_df["cob_meses"] < u_cob)].sort_values("conv_vcto_dias")
    sin_conv_riesgo = df[df["convenio"].isna() & df["estado"].isin(["QS", "CR"]) & (df["consumo_prom"] > 0)]
    d1, d2 = st.columns(2)
    with d1:
        kpi("Convenios en alerta", n0(len(al)), f"De {n0(len(cv_df))} fármacos con convenio", "#E8890C")
    with d2:
        kpi("QS o CR sin convenio vigente", n0(len(sin_conv_riesgo)), "Dependen de compra directa o Cenabast", "#B3261E")
    st.dataframe(
        al[["estado_icon", "descripcion", "convenio", "conv_saldo", "cob_meses", "conv_vcto", "conv_vcto_dias"]],
        hide_index=True,
        column_config={
            "estado_icon": "Estado", "descripcion": "Fármaco", "convenio": "Licitación",
            "conv_saldo": st.column_config.NumberColumn("Saldo", format="%.0f"),
            "cob_meses": st.column_config.NumberColumn("Cobertura del saldo (meses)", format="%.1f"),
            "conv_vcto": st.column_config.DateColumn("Vence", format="DD/MM/YYYY"),
            "conv_vcto_dias": st.column_config.NumberColumn("Días al vencimiento", format="%.0f")})

    # 5. Vencimientos de lotes de los FÁRMACOS DE MAYOR DEMANDA (¡Nuevo y detallado igual que insumos!)
    st.subheader("Vencimientos por lote (Fármacos de mayor demanda)")
    if archivos["vencimientos"] is None and venc.empty:
        st.info("Sube el reporte de stock por lote para ver qué fármacos de mayor demanda están por vencer.")
    else:
        try:
            if not venc.empty:
                vl_f = venc.copy()
                meta_vf = {"fecha": fecha_stock, "filtro": "Próximos meses"}
            else:
                vl_f, meta_vf = load_vencimientos(archivos["vencimientos"][1])
        except Exception as exc:
            st.error(f"No pude leer el reporte de vencimientos: {exc}")
        else:
            ref_vf = pd.Timestamp(meta_vf["fecha"] or date.today())
            # Renombrar columnas si vienen del archivo crudo
            if "Código" in vl_f.columns:
                vl_f = vl_f.rename(columns={"Código": "codigo", "Fecha Vencimiento": "vence", "Cantidad": "cantidad", "Precio Lote": "precio_lote"})
            if "Fecha Vencimiento" in vl_f.columns:
                vl_f["vence"] = pd.to_datetime(vl_f["Fecha Vencimiento"], errors="coerce")
            
            lots_f = riesgo_fefo(
                vl_f[vl_f["codigo"].isin(set(b["codigo"]))],
                b.set_index("codigo")["tasa_fc_dia"], 
                ref_vf
            )
            # Preparamos un DataFrame de fármacos renombrando 'descripcion' a 'producto' para render_venc_top
            prod_farm = df.rename(columns={"descripcion": "producto"}).copy()
            render_venc_top(lots_f, meta_vf, prod_farm, "farm", "fármacos", ref_vf)

with tab5:
    render_insumos(archivos["stock_insumos"], archivos["salidas_insumos"], data["licit"],
                   u_cr, u_ss, horiz_m, dias_seg, archivos["vencimientos"])
